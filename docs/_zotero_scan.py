# -*- coding: utf-8 -*-
"""Read Zotero change-detection related papers + extract discussion docx."""
import shutil
import sqlite3
import tempfile
from pathlib import Path

from docx import Document

DB = Path(r"C:\Users\Ruiku\Zotero\zotero.sqlite")
DOCX = Path(r"C:\Users\Ruiku\WPSDrive\382620557\WPS云盘\博士\战争场景变化检测\沟通记录20260918.docx")

print("=" * 60)
print("DISCUSSION RECORD")
print("=" * 60)
doc = Document(str(DOCX))
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        print(t)
for ti, table in enumerate(doc.tables):
    print(f"\n--- table {ti} ---")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        print(" | ".join(cells))

print("\n" + "=" * 60)
print("ZOTERO")
print("=" * 60)

tmp = Path(tempfile.gettempdir()) / "zotero_ro_copy.sqlite"
shutil.copy2(DB, tmp)
con = sqlite3.connect(str(tmp))
cur = con.cursor()

cur.execute(
    """
    SELECT c.collectionID, c.collectionName
    FROM collections c
    ORDER BY c.collectionName
    """
)
print("\n--- collections ---")
collections = cur.fetchall()
for cid, name in collections:
    print(f"{cid}\t{name}")

# Find CD-related collections
keywords = ["变化", "change", "CD", "检测", "遥感", "损毁", "damage", "semantic", "SCD", "样本", "domain"]
cd_cids = []
for cid, name in collections:
    low = (name or "").lower()
    if any(k.lower() in low or k in (name or "") for k in keywords):
        cd_cids.append((cid, name))

print("\n--- CD-related collections ---")
for cid, name in cd_cids:
    print(f"{cid}\t{name}")

# Items in those collections with titles
sql = """
SELECT DISTINCT
  c.collectionName,
  idv.value AS title,
  idv2.value AS year,
  idv3.value AS pub,
  i.itemID
FROM collectionItems ci
JOIN collections c ON c.collectionID = ci.collectionID
JOIN items i ON i.itemID = ci.itemID
JOIN itemData id ON id.itemID = i.itemID AND id.fieldID = (
  SELECT fieldID FROM fields WHERE fieldName='title'
)
JOIN itemDataValues idv ON idv.valueID = id.valueID
LEFT JOIN itemData id2 ON id2.itemID = i.itemID AND id2.fieldID = (
  SELECT fieldID FROM fields WHERE fieldName='date'
)
LEFT JOIN itemDataValues idv2 ON idv2.valueID = id2.valueID
LEFT JOIN itemData id3 ON id3.itemID = i.itemID AND id3.fieldID = (
  SELECT fieldID FROM fields WHERE fieldName='publicationTitle'
)
LEFT JOIN itemDataValues idv3 ON idv3.valueID = id3.valueID
WHERE ci.collectionID IN ({})
  AND idv.value IS NOT NULL
ORDER BY c.collectionName, idv.value
""".format(
    ",".join(str(c) for c, _ in cd_cids) if cd_cids else "-1"
)

print("\n--- items in CD collections ---")
try:
    cur.execute(sql)
    rows = cur.fetchall()
    for r in rows:
        print("\t".join(str(x) if x is not None else "" for x in r[:4]))
    print(f"\nTOTAL: {len(rows)}")
except Exception as e:
    print("SQL error:", e)
    # fallback: search titles
    cur.execute(
        """
        SELECT idv.value, idv2.value
        FROM items i
        JOIN itemData id ON id.itemID=i.itemID AND id.fieldID=(SELECT fieldID FROM fields WHERE fieldName='title')
        JOIN itemDataValues idv ON idv.valueID=id.valueID
        LEFT JOIN itemData id2 ON id2.itemID=i.itemID AND id2.fieldID=(SELECT fieldID FROM fields WHERE fieldName='date')
        LEFT JOIN itemDataValues idv2 ON idv2.valueID=id2.valueID
        WHERE lower(idv.value) LIKE '%change%'
           OR lower(idv.value) LIKE '%damage%'
           OR idv.value LIKE '%变化%'
           OR lower(idv.value) LIKE '%semantic change%'
           OR lower(idv.value) LIKE '%war%'
           OR lower(idv.value) LIKE '%remote sensing%'
        ORDER BY idv.value
        LIMIT 200
        """
    )
    for r in cur.fetchall():
        print("\t".join(str(x) if x else "" for x in r))
    print("TOTAL fallback:", cur.execute("SELECT 1").fetchone())

# Also dump ALL collection names with item counts
print("\n--- all collection item counts ---")
cur.execute(
    """
    SELECT c.collectionName, COUNT(*) AS n
    FROM collectionItems ci
    JOIN collections c ON c.collectionID=ci.collectionID
    GROUP BY c.collectionName
    ORDER BY n DESC
    """
)
for name, n in cur.fetchall():
    print(f"{n:4d}  {name}")

con.close()
print("\nDONE")
