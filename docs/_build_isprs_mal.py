# -*- coding: utf-8 -*-
"""组会精读 PPT v3：模块2只贴图；加重马磊关注点。"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(r"F:\mambacd\home\HICD_v7\docs\ISPRS2024_精读_王瑞坤.pptx")
FIG = Path(r"F:\mambacd\home\HICD_v7\docs\assets\paper_figs")
F1, F2, F3, F4 = FIG / "fig1_workflow.png", FIG / "fig2_mpanet.png", FIG / "fig3_tc.png", FIG / "fig4_sdi.png"
F5, F6 = FIG / "fig5_datasets.png", FIG / "fig6_ablation_qual.png"
T2 = FIG / "table2_ablation.png"

PURPLE = RGBColor(0x6E, 0x0F, 0x6D)
RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x0D, 0x73, 0x7A)
ORANGE = RGBColor(0xC4, 0x5C, 0x26)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_PURPLE = RGBColor(0xF3, 0xE8, 0xF3)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xF8)
LIGHT_RED = RGBColor(0xFB, 0xEB, 0xEB)
LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
LIGHT_YELLOW = RGBColor(0xFF, 0xF8, 0xE1)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)

FONT = "微软雅黑"
L = Emu(280000)
FULL_W = Emu(8500000)


def set_cjk(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    succ = {
        "a:ea": ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
        "a:cs": ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
    }
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, *succ[tag])
        el.set("typeface", name)


def style(run, size=14, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_cjk(run, name)


def text(slide, x, y, w, h, segs, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    if isinstance(segs, str):
        segs = [(segs, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    for t, kw in segs:
        if t == "\n":
            p = tf.add_paragraph()
            p.alignment = kw.get("align", align)
            continue
        r = p.add_run()
        r.text = t
        style(r, size=kw.get("size", 14), bold=kw.get("bold", False),
              color=kw.get("color", INK), name=kw.get("font", FONT))
    return box


def list_card(slide, x, y, w, h, title, items, *, fill=LIGHT_BLUE, title_color=BLUE, body_size=13):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = title_color
    sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    text(slide, Emu(x + 100000), Emu(y + 50000), Emu(w - 200000), Emu(250000),
         [(title, {"size": 13, "bold": True, "color": title_color})])
    box = slide.shapes.add_textbox(Emu(x + 100000), Emu(y + 320000), Emu(w - 200000), Emu(h - 380000))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = "• " + it
        style(r, size=body_size, color=INK)


def header(slide, page=None):
    text(slide, L, Emu(80000), Emu(1600000), Emu(320000),
         [("论文阅读", {"size": 18, "bold": True, "color": PURPLE})])
    if page:
        text(slide, Emu(8300000), Emu(7000000), Emu(500000), Emu(220000),
             [(str(page), {"size": 10, "color": GRAY})], align=PP_ALIGN.RIGHT)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, Emu(420000), FULL_W, Emu(14000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PURPLE
    bar.line.fill.background()


def title(slide, s, y=Emu(520000)):
    text(slide, L, y, FULL_W, Emu(360000), [(s, {"size": 20, "bold": True, "color": PURPLE})])


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def img(slide, path, x, y, w, h=None):
    p = Path(path)
    if not p.exists():
        text(slide, x, y, w, Emu(300000), [(f"[缺图 {p.name}]", {"size": 12, "color": RED})])
        return
    if h is None:
        slide.shapes.add_picture(str(p), x, y, width=w)
    else:
        slide.shapes.add_picture(str(p), x, y, width=w, height=h)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 1 封面
    s = new(prs)
    text(s, L, Emu(1800000), FULL_W, Emu(400000),
         [("文献精读 · 组会", {"size": 16, "bold": True, "color": PURPLE})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(2300000), FULL_W, Emu(1000000),
         [("HGINet：层级语义图交互网络\n用于高分辨率遥感语义变化检测",
           {"size": 24, "bold": True})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(3600000), FULL_W, Emu(700000),
         [("Long et al., ISPRS J. Photogramm. Remote Sens., 2024",
           {"size": 14, "color": GRAY})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(4200000), FULL_W, Emu(400000),
         [("汇报人：王瑞坤", {"size": 16, "bold": True})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(4700000), FULL_W, Emu(350000),
         [("重点：图建模关系 — 和我们战争场景变化检测的接口",
           {"size": 13, "bold": True, "color": RED})], align=PP_ALIGN.CENTER)
    notes(s, "组会精读；突出马磊关心的关系建模与任务设定。")

    # 2 问题
    s = new(prs)
    header(s, 2)
    title(s, "要解决什么问题？")
    text(s, L, Emu(1000000), FULL_W, Emu(450000),
         [("SCD：定位「哪里变」+ 判断「变成什么」。作者认为关键不在堆网络，而在建模关系。",
           {"size": 15, "bold": True})])
    list_card(s, L, Emu(1600000), Emu(4000000), Emu(2100000), "难点", [
        "变化区域小、背景占主导",
        "复杂场景易假变化",
        "双时相关系常被忽略",
    ], fill=LIGHT_RED, title_color=RED)
    list_card(s, Emu(4600000), Emu(1600000), Emu(4150000), Emu(2100000), "三步思路", [
        "MPANet 提特征",
        "TC 稳住未变化",
        "SDI 用 GCN 做交互",
    ], fill=LIGHT_GREEN, title_color=TEAL)
    text(s, L, Emu(4000000), FULL_W, Emu(1200000),
         [("和我们场景的对照（先记下）：",
           {"size": 14, "bold": True, "color": RED}),
          ("\n他们是城市地表 from→to；我们是特殊目标 + 损毁状态，且目标更稀疏、更不均衡。",
            {"size": 14})])
    notes(s, "问题页埋下与战争场景差异，后面重点展开。")

    # 3 总框架
    s = new(prs)
    header(s, 3)
    title(s, "方法总览（论文 Fig.1）")
    img(s, F1, L, Emu(950000), FULL_W, Emu(2800000))
    text(s, L, Emu(3850000), FULL_W, Emu(220000),
         [("Fig.1  MPANet → TC → SDI", {"size": 10, "color": GRAY})])
    text(s, L, Emu(4200000), FULL_W, Emu(1200000),
         [("老师可能关心的一点：",
           {"size": 14, "bold": True, "color": RED}),
          ("\n这里的「图」建在特征层交互上，不是建在目标/标签关系上——后面会展开。",
            {"size": 15, "bold": True})])
    notes(s, "总图之后立刻点出：图建在哪一层。")

    # 4 MPANet
    s = new(prs)
    header(s, 4)
    title(s, "模块① MPANet：特征提取（Fig.2）")
    img(s, F2, L, Emu(900000), Emu(5400000), Emu(4400000))
    text(s, L, Emu(5400000), Emu(5400000), Emu(200000),
         [("Fig.2", {"size": 10, "color": GRAY})])
    list_card(s, Emu(5700000), Emu(1000000), Emu(3050000), Emu(2500000), "一句话", [
        "多尺度 Transformer 特征",
        "远程+近程上下文",
        "边界细节 refine",
    ], fill=LIGHT_BLUE, title_color=BLUE)
    text(s, Emu(5700000), Emu(3700000), Emu(3050000), Emu(1800000),
         [("对战争场景：\n线状目标（跑道等）或可借条带注意力思路。",
           {"size": 13, "bold": True})])
    notes(s, "MPANet从简，一笔带过即可。")

    # 5 TC —— 只贴图，不写原文段落
    s = new(prs)
    header(s, 5)
    title(s, "模块② TC：时序相关（Fig.3）")
    # 图居中放大，仅一行说明
    img(s, F3, Emu(600000), Emu(1100000), Emu(8000000), Emu(4200000))
    text(s, L, Emu(5600000), FULL_W, Emu(500000),
         [("Fig.3  双时相交叉注意力：让「没变」的区域特征更一致",
           {"size": 14, "bold": True, "color": PURPLE})], align=PP_ALIGN.CENTER)
    notes(s, "本页只展示论文Fig.3，口头一句话即可，不展开原文。")

    # 6 SDI —— 马磊最可能关心
    s = new(prs)
    header(s, 6)
    title(s, "模块③ SDI：图卷积交互（Fig.4）★ 老师可能最关心")
    img(s, F4, L, Emu(900000), FULL_W, Emu(3300000))
    text(s, L, Emu(4250000), FULL_W, Emu(200000),
         [("Fig.4  多层特征 → 图节点 → GCN 交互 → 变化特征", {"size": 10, "color": GRAY})])
    list_card(s, L, Emu(4550000), Emu(4000000), Emu(2000000), "为什么可能感兴趣", [
        "图适合不规则/多层关系",
        "消融显示 SDI 增益最大",
        "「关系建模」有学术抓手",
    ], fill=LIGHT_YELLOW, title_color=ORANGE)
    list_card(s, Emu(4600000), Emu(4550000), Emu(4150000), Emu(2000000), "要追问的差异", [
        "图在特征层，不在标签层",
        "from→to ≠ 目标×状态",
        "没有实例清单/可清点输出",
    ], fill=LIGHT_RED, title_color=RED)
    notes(s, "SDI是重点页：肯定GCN价值，同时指出与我们标签图GNN的差别。")

    # 7 马磊兴趣点专页
    s = new(prs)
    header(s, 7)
    title(s, "这篇文章里，老师可能感兴趣的点")
    # 三列重点卡
    list_card(s, L, Emu(1000000), Emu(2700000), Emu(3000000), "① 图 / 关系建模", [
        "GCN 处理不规则结构",
        "与我们 GNN 标签空间同属「关系」路线",
        "可讨论：关系放特征还是放标签",
    ], fill=LIGHT_PURPLE, title_color=PURPLE, body_size=12)
    list_card(s, Emu(3050000), Emu(1000000), Emu(2700000), Emu(3000000), "② 任务设定边界", [
        "标准 SCD：地类 from→to",
        "战争场景：特殊目标+损毁态",
        "方法可借鉴，任务不能照搬",
    ], fill=LIGHT_BLUE, title_color=BLUE, body_size=12)
    list_card(s, Emu(5800000), Emu(1000000), Emu(2900000), Emu(3000000), "③ 稀疏与未变化", [
        "变化少、背景多是共同现实",
        "TC 强化未变化有启发",
        "我们长尾更极端（弹坑等）",
    ], fill=LIGHT_GREEN, title_color=TEAL, body_size=12)

    list_card(s, L, Emu(4200000), Emu(4000000), Emu(2200000), "④ 输出形态", [
        "本文：分割/变化图",
        "应用制图还需要：目标清点、状态统计",
        "实例级/清单式输出是缺口",
    ], fill=LIGHT_RED, title_color=RED, body_size=12)
    list_card(s, Emu(4600000), Emu(4200000), Emu(4150000), Emu(2200000), "⑤ 组会可讨论", [
        "弹坑/战壕：单点还是成群区域？",
        "统一框架 vs 分目标不同方法？",
        "样本与结构，战争场景谁更关键？",
    ], fill=LIGHT_YELLOW, title_color=ORANGE, body_size=12)
    notes(s, "这页是给马磊的：图关系、任务边界、稀疏、输出形态、讨论题。")

    # 8 数据
    s = new(prs)
    header(s, 8)
    title(s, "数据与指标（Fig.5）")
    img(s, F5, L, Emu(950000), FULL_W, Emu(3200000))
    text(s, L, Emu(4250000), FULL_W, Emu(200000),
         [("Fig.5  福州 / 厦门双时相样例", {"size": 10, "color": GRAY})])
    list_card(s, L, Emu(4550000), Emu(4000000), Emu(1700000), "实验设置", [
        "SECOND、HRSCD + 福州、厦门",
        "指标 mIoU / SeK / F1scd",
    ], fill=LIGHT_GRAY, title_color=GRAY)
    text(s, Emu(4600000), Emu(4650000), Emu(4100000), Emu(1400000),
         [("公开集标签噪声大——\n与我们「标注质量参差」是同类问题。",
           {"size": 13, "bold": True})])
    notes(s, "数据从简。")

    # 9 结果
    s = new(prs)
    header(s, 9)
    title(s, "效果（主结果）")
    cats = ["SECOND", "HRSCD", "Fuzhou", "Xiamen"]
    vals = [59.48, 64.12, 64.45, 84.93]
    chart_x, chart_y = L, Emu(1000000)
    chart_w, chart_h = Emu(5200000), Emu(3000000)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, chart_x, chart_y, chart_w, chart_h)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xF7, 0xF7, 0xFA)
    bg.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    text(s, chart_x, Emu(chart_y - 260000), chart_w, Emu(240000),
         [("F1scd (%)", {"size": 12, "bold": True, "color": PURPLE})])
    n = len(cats)
    gw = chart_w / n
    for i, (c, v) in enumerate(zip(cats, vals)):
        bh = int(chart_h * 0.82 * (v / 100.0))
        bx = Emu(chart_x + i * gw + int(gw * 0.25))
        bw = Emu(int(gw * 0.5))
        by = Emu(chart_y + chart_h - 180000 - bh)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, by, bw, Emu(bh))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PURPLE
        bar.line.fill.background()
        text(s, Emu(bx - 60000), Emu(by - 210000), Emu(bw + 120000), Emu(190000),
             [(f"{v}", {"size": 11, "bold": True, "color": PURPLE})], align=PP_ALIGN.CENTER)
        text(s, Emu(chart_x + i * gw), Emu(chart_y + chart_h + 20000), Emu(gw), Emu(200000),
             [(c, {"size": 11})], align=PP_ALIGN.CENTER)
    list_card(s, Emu(5700000), Emu(1000000), Emu(3050000), Emu(3000000), "读结果时想到的", [
        "九基线中总体领先",
        "厦门高 ≠ 任务简单",
        "复杂集上更有说服力",
        "但仍是 from→to 任务成绩",
    ], fill=LIGHT_BLUE, title_color=BLUE, body_size=12)
    text(s, L, Emu(4500000), FULL_W, Emu(800000),
         [("提醒：这套指标证明的是「城市语义变化」做得好，不能直接外推到战争特殊目标。",
           {"size": 14, "bold": True, "color": RED})])
    notes(s, "结果之后立刻加边界提醒，服务马磊讨论。")

    # 10 消融
    s = new(prs)
    header(s, 10)
    title(s, "消融：论文原表 + 原图")
    img(s, T2, L, Emu(900000), Emu(7200000), Emu(1900000))
    text(s, L, Emu(2850000), Emu(7200000), Emu(200000),
         [("Table 2  完整模型 F1scd 59.48；SDI 单独增益最明显", {"size": 10, "color": GRAY})])
    img(s, F6, L, Emu(3150000), Emu(7000000), Emu(2400000))
    text(s, L, Emu(5600000), FULL_W, Emu(400000),
         [("Fig.6  定性对比：加模块后边界与类别更接近真值",
           {"size": 12, "bold": True})])
    notes(s, "用原表原图；口头强调SDI。")

    # 11 启发 + 文献
    s = new(prs)
    header(s, 11)
    title(s, "小结与文献")
    list_card(s, L, Emu(1000000), Emu(4000000), Emu(2400000), "三句话带走", [
        "关系建模有价值（图/时序）",
        "但任务设定是城市 from→to",
        "战争场景还要目标×状态+可清点",
    ], fill=LIGHT_PURPLE, title_color=PURPLE)
    list_card(s, Emu(4600000), Emu(1000000), Emu(4150000), Emu(2400000), "建议讨论", [
        "图：特征层 vs 标签层",
        "稀有特殊目标如何检测",
        "统一框架是否必要",
    ], fill=LIGHT_YELLOW, title_color=ORANGE)
    text(s, L, Emu(3700000), FULL_W, Emu(1200000),
         [("Long et al., 2024. ISPRS J. 211, 318–335.\nDOI: 10.1016/j.isprsjprs.2024.04.012",
           {"size": 13, "color": GRAY, "font": "Times New Roman"})])
    text(s, L, Emu(5200000), FULL_W, Emu(400000),
         [("谢谢，请老师批评指正", {"size": 18, "bold": True, "color": PURPLE})],
         align=PP_ALIGN.CENTER)
    notes(s, "收尾三句话对齐马磊关注点。")

    prs.save(str(OUT))
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    n = sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    print("SAVED", OUT, "slides", len(prs.slides), "pictures", n)


if __name__ == "__main__":
    main()
