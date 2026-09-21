# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu

path = r"C:\Users\Ruiku\WPSDrive\382620557\WPS云盘\review\周报PPT\王瑞琨20231124进展.pptx"
prs = Presentation(path)
print("slide_size_emu", prs.slide_width, prs.slide_height)
print("slide_size_in", prs.slide_width / 914400, prs.slide_height / 914400)
print("slide_count", len(prs.slides))
print("layouts", [l.name for l in prs.slide_masters[0].slide_layouts])

for i, slide in enumerate(prs.slides):
    print("===SLIDE", i + 1, "layout=", slide.slide_layout.name)
    for shape in slide.shapes:
        t = ""
        if shape.has_text_frame:
            t = " | ".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        print(
            f"  shape name={shape.name!r} type={shape.shape_type} "
            f"l={shape.left} t={shape.top} w={shape.width} h={shape.height}"
        )
        if t:
            print("    text:", t[:800])
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    f = r.font
                    try:
                        color = f.color.rgb if f.color and f.color.type is not None else None
                    except Exception:
                        color = None
                    print(
                        f"    run: {r.text[:100]!r} size={f.size} bold={f.bold} "
                        f"name={f.name} color={color}"
                    )
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                print("  NOTES:", notes[:400])
    except Exception as e:
        print("  notes err", e)
