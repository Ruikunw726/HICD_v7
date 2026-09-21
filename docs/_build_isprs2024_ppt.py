# -*- coding: utf-8 -*-
"""Build ISPRS 2024 HGINet paper-reading PPT in weekly-report style."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(r"F:\mambacd\home\HICD_v7\docs\ISPRS2024_精读_王瑞坤.pptx")
FLOW = Path(r"F:\mambacd\home\HICD_v7\docs\assets\hginet_flowchart.png")
FLOW_SMALL = Path(r"F:\mambacd\home\HICD_v7\docs\assets\hginet_flowchart_small.png")
ASSETS = Path(r"F:\mambacd\home\HICD_v7\docs\assets")
ASSETS.mkdir(parents=True, exist_ok=True)

PURPLE = RGBColor(0x6E, 0x0F, 0x6D)
RED = RGBColor(0xFF, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x00, 0x00, 0xFF)
DARK_RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_PURPLE = RGBColor(0xF3, 0xE8, 0xF3)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xF8)
LIGHT_RED = RGBColor(0xFB, 0xEB, 0xEB)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)

FONT = "微软雅黑"
FONT_EN = "Times New Roman"

SW = Inches(10)
SH = Inches(7.5)
# Safe content band: x in [L, R]
L = Emu(280000)
R = Emu(8864000)  # 8864000
FULL_W = Emu(R - L)  # 8584000
HALF_W = Emu(4050000)
RIGHT_X = Emu(4600000)


def set_cjk_font(run, font_name=FONT):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    successors = {
        "a:ea": ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
        "a:cs": ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
    }
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, *successors[tag])
        el.set("typeface", font_name)


def style_run(run, size=16, bold=False, color=INK, font_name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_cjk_font(run, font_name)


def add_text(slide, left, top, width, height, segments, *,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             word_wrap=True, fill=None, line=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = valign
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        if line is None:
            box.line.fill.background()
    if line is not None:
        box.line.color.rgb = line
    if isinstance(segments, str):
        segments = [{"text": segments}]
    if segments and isinstance(segments[0], list):
        paras = segments
    else:
        paras = [segments]
    first = True
    for items in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if isinstance(items, str):
            items = [{"text": items}]
        for item in items:
            if isinstance(item, str):
                item = {"text": item}
            run = p.add_run()
            run.text = item.get("text", "")
            style_run(run, size=item.get("size", 16), bold=item.get("bold", False),
                      color=item.get("color", INK), font_name=item.get("font", FONT))
    return box


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def add_line(slide, left, top, width, color=PURPLE, weight=1.25):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(weight))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_content_header(slide, page_no, title="论文阅读"):
    add_text(slide, Emu(0), Emu(100000), SW, Emu(640000),
             [{"text": title, "size": 28, "color": PURPLE}], align=PP_ALIGN.LEFT)
    add_line(slide, Emu(0), Emu(765000), SW, color=PURPLE, weight=1.2)
    add_text(slide, Emu(7800000), Emu(6650000), Emu(1200000), Emu(280000),
             [{"text": str(page_no), "size": 12, "color": GRAY}], align=PP_ALIGN.RIGHT)


def add_red_label(slide, text, top=Emu(800000)):
    return add_text(slide, L, top, FULL_W, Emu(380000),
                    [{"text": text, "size": 18, "color": RED}],
                    valign=MSO_ANCHOR.MIDDLE)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def make_flow_small():
    if not FLOW.exists():
        return None
    img = Image.open(FLOW)
    img.thumbnail((1600, 700), Image.Resampling.LANCZOS)
    img.save(FLOW_SMALL, optimize=True)
    return FLOW_SMALL


def add_picture_fit(slide, path, left, top, max_w, max_h):
    if not path or not Path(path).exists():
        return None
    img = Image.open(path)
    iw, ih = img.size
    aspect = iw / ih
    w = max_w
    h = int(w / aspect)
    if h > max_h:
        h = max_h
        w = int(h * aspect)
    return slide.shapes.add_picture(str(path), left, top, width=int(w), height=int(h))


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    flow = make_flow_small()

    # ---------- 1 Cover ----------
    s = blank_slide(prs)
    t1 = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Emu(-20000), Emu(-30000), Emu(2200000), Emu(1800000))
    t1.fill.solid(); t1.fill.fore_color.rgb = PURPLE; t1.line.fill.background()
    t2 = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Emu(7200000), Emu(1200000), Emu(2200000), Emu(2000000))
    t2.fill.solid(); t2.fill.fore_color.rgb = RGBColor(0x9A, 0x2F, 0x98); t2.line.fill.background()
    t3 = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Emu(7800000), Emu(5200000), Emu(1600000), Emu(2000000))
    t3.fill.solid(); t3.fill.fore_color.rgb = LIGHT_PURPLE; t3.line.fill.background(); t3.rotation = 180
    add_rect(s, Emu(400000), Emu(2500000), Emu(8200000), Emu(2000000), LIGHT_PURPLE)
    add_text(s, Emu(700000), Emu(2700000), Emu(7800000), Emu(700000),
             [{"text": "文献精读报告", "size": 20, "color": PURPLE, "bold": True}], align=PP_ALIGN.CENTER)
    add_text(s, Emu(500000), Emu(3250000), Emu(8200000), Emu(1000000),
             [
                 {"text": "Semantic Change Detection using a Hierarchical\n", "size": 20, "color": INK, "bold": True, "font": FONT_EN},
                 {"text": "Semantic Graph Interaction Network (HGINet)", "size": 20, "color": PURPLE, "bold": True, "font": FONT_EN},
             ], align=PP_ALIGN.CENTER)
    add_text(s, Emu(700000), Emu(4300000), Emu(7800000), Emu(400000),
             [{"text": "ISPRS J. Photogramm. Remote Sens., 2024, 211: 318–335", "size": 14, "color": GRAY, "font": FONT_EN}],
             align=PP_ALIGN.CENTER)
    add_text(s, Emu(4800000), Emu(5450000), Emu(3800000), Emu(800000),
             [
                 [{"text": "汇报人：王瑞坤", "size": 16, "color": INK}],
                 [{"text": "对应课题：HICD 层级语义变化检测", "size": 14, "color": GRAY}],
             ], align=PP_ALIGN.RIGHT)
    add_text(s, Emu(500000), Emu(5700000), Emu(4200000), Emu(500000),
             [
                 [{"text": "Zotero Key：WY3KTIRY", "size": 13, "color": GRAY}],
                 [{"text": "DOI：10.1016/j.isprsjprs.2024.04.012", "size": 12, "color": GRAY, "font": FONT_EN}],
             ])
    add_notes(s, "精读对象为 ISPRS 2024 HGINet，与博士课题层级语义变化检测直接相关。")

    # ---------- 2 Agenda ----------
    s = blank_slide(prs)
    add_rect(s, Emu(0), Emu(250000), SW, Emu(1650000), LIGHT_PURPLE)
    add_text(s, Emu(0), Emu(700000), SW, Emu(700000),
             [{"text": "汇报大纲", "size": 36, "color": PURPLE}], align=PP_ALIGN.CENTER)
    add_line(s, Emu(350000), Emu(2100000), Emu(8400000), color=PURPLE)
    items = [
        ("01", "选文理由", "顶刊 SCD + 图交互 + 与 HICD 同构"),
        ("02", "研究背景与问题", "时序相关 vs 变化语义耦合"),
        ("03", "方法结构拆解", "MPA / TC / SDI 三模块"),
        ("04", "数据集与实验", "四数据集 F1scd / SeK"),
        ("05", "对照 HICD 启示", "特征图 GCN vs 标签图 GNN"),
        ("06", "挑战与行动项", "基线对比与消融落地"),
    ]
    y = 2400000
    for num, title, sub in items:
        add_rect(s, Emu(700000), Emu(y), Emu(800000), Emu(520000), PURPLE)
        add_text(s, Emu(700000), Emu(y), Emu(800000), Emu(520000),
                 [{"text": num, "size": 16, "color": WHITE, "bold": True, "font": FONT_EN}],
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(s, Emu(1700000), Emu(y), Emu(6800000), Emu(520000),
                 [
                     {"text": title + "　", "size": 18, "color": INK, "bold": True},
                     {"text": sub, "size": 14, "color": GRAY},
                 ], valign=MSO_ANCHOR.MIDDLE)
        y += 620000
    add_text(s, Emu(7800000), Emu(6650000), Emu(1200000), Emu(280000),
             [{"text": "2", "size": 12, "color": GRAY}], align=PP_ALIGN.RIGHT)
    add_notes(s, "按选文—背景—方法—实验—对照—行动六段推进。")

    # ---------- 3 Why this paper ----------
    s = blank_slide(prs)
    add_content_header(s, 3)
    add_red_label(s, "选文理由：", top=Emu(820000))
    add_rect(s, L, Emu(1250000), FULL_W, Emu(1050000), LIGHT_GRAY)
    add_text(s, Emu(380000), Emu(1320000), Emu(8300000), Emu(900000),
             [
                 [
                     {"text": "Long J., Li M., Wang X., Stein A. ", "size": 13, "font": FONT_EN},
                     {"text": "Semantic change detection using a hierarchical semantic graph interaction network…", "size": 13, "color": PURPLE, "font": FONT_EN},
                 ],
                 [{"text": "ISPRS J. Photogramm. Remote Sens., 2024, 211: 318–335. DOI: 10.1016/j.isprsjprs.2024.04.012", "size": 12, "color": GRAY, "font": FONT_EN}],
                 [{"text": "被引：WoS 核心约 88 / 总约 91；代码：github.com/long123524/HGINet-torch", "size": 12, "color": GRAY}],
             ])
    reasons = [
        ("① 直接对标博士问题", "任务即 hierarchical / semantic change detection，与 HICD「目标×状态」同属 SCD 族；P1 目标期刊含 ISPRS，同赛道代表作。", BLUE),
        ("② 图交互与标签图互补", "HGINet 把 GCN 用在双时相特征与差特征交互；HICD 的 GNNLabelStructure 把图用在 15 目标 + 7 状态标签拓扑。", DARK_RED),
        ("③ 协议与基线可复用", "报告 SECOND / HRSCD / Fuzhou / Xiamen 的 F1scd、SeK，并与 9 个方法比较，覆盖 baseline_compare_plan 缺口。", BLUE),
        ("④ 可复现与可扩展", "官方开源 + PVT-v2-B2 权重；同组 2025 续作 BGSNet 可作为跟踪线索。", BLUE),
    ]
    y = 2450000
    for title, body, color in reasons:
        add_text(s, L, Emu(y), FULL_W, Emu(850000),
                 [
                     [{"text": title, "size": 15, "bold": True, "color": color}],
                     [{"text": body, "size": 12, "color": INK}],
                 ])
        y += 900000
    add_notes(s, "选文四条：同赛道顶刊、图用法互补、协议/基线可复用、开源可复现。")

    # ---------- 4 Background ----------
    s = blank_slide(prs)
    add_content_header(s, 4)
    add_red_label(s, "研究背景：", top=Emu(820000))
    add_text(s, L, Emu(1250000), FULL_W, Emu(750000),
             [
                 [{"text": "语义变化检测（SCD）不仅要回答「哪里变了」，还要回答「从哪类变成哪类」。", "size": 15, "bold": True}],
                 [{"text": "难点：双时相语义特征、差分特征及二者时序相关（TC）建模不足，复杂场景下变化区域与类别均易出错。", "size": 13}],
             ])
    cards = [
        ("痛点 1", "时序相关弱", "未变化区缺少跨时相一致性，易产生季节/光照伪变化。", LIGHT_RED),
        ("痛点 2", "差分语义浅", "简单绝对差难以表达谁参与变化、如何交互。", LIGHT_BLUE),
        ("痛点 3", "任务割裂", "分割与变化若分治，难共享层级语义上下文。", LIGHT_PURPLE),
    ]
    x = 280000
    cw = 2700000
    for title, head, body, fill in cards:
        add_rect(s, Emu(x), Emu(2300000), Emu(cw), Emu(2100000), fill)
        add_text(s, Emu(x + 100000), Emu(2400000), Emu(cw - 200000), Emu(1900000),
                 [
                     [{"text": title, "size": 12, "color": RED, "bold": True}],
                     [{"text": head, "size": 15, "color": PURPLE, "bold": True}],
                     [{"text": body, "size": 12, "color": INK}],
                     [{"text": "→ 对应 HICD 设计动机", "size": 11, "color": GRAY}],
                 ])
        x += 2850000
    add_text(s, L, Emu(4600000), FULL_W, Emu(1600000),
             [
                 [{"text": "作者目标", "size": 14, "color": RED, "bold": True}],
                 [{"text": "构造多任务网络 HGINet，同时提升「变化区域」与「变化类别」识别，服务复杂 SCD 场景。", "size": 13}],
                 [{"text": "一句话：", "size": 13, "color": DARK_RED, "bold": True},
                  {"text": "层级特征提取 + 跨时相相关增强 + 差异语义图交互，拧成一条链。", "size": 13, "color": DARK_RED}],
                 [{"text": "与伪变化的关系：TC 强调未变区一致性，与我们域差距/季节伪变化问题同构。", "size": 12, "color": BLUE}],
             ])
    add_notes(s, "背景强调 TC 建模不足与伪变化问题，对应 HICD 中季节/域差距伪变化。")

    # ---------- 5 Problem definition ----------
    s = blank_slide(prs)
    add_content_header(s, 5)
    add_red_label(s, "问题定义与任务形式：", top=Emu(820000))
    add_rect(s, L, Emu(1250000), Emu(4000000), Emu(3600000), LIGHT_GRAY)
    add_text(s, Emu(380000), Emu(1350000), Emu(3800000), Emu(3400000),
             [
                 [{"text": "输入", "size": 14, "color": RED, "bold": True}],
                 [{"text": "双时相高分影像 I¹, I²（及标签 L¹, L²）", "size": 12}],
                 [{"text": "输出（多任务）", "size": 14, "color": RED, "bold": True}],
                 [{"text": "① 变化图 / 变化类别图（SCD）", "size": 12, "color": BLUE}],
                 [{"text": "② 时相 1 / 时相 2 语义分割图", "size": 12, "color": BLUE}],
                 [{"text": "监督信号", "size": 14, "color": RED, "bold": True}],
                 [{"text": "像元级 from-to 语义变化标注", "size": 12}],
                 [{"text": "评价", "size": 14, "color": RED, "bold": True}],
                 [{"text": "F1scd + SeK（SCD 标准协议）", "size": 12}],
                 [{"text": "标签组织", "size": 14, "color": RED, "bold": True}],
                 [{"text": "类别集合近似对等，非目标⊕状态层级", "size": 12, "color": DARK_RED}],
             ])
    add_text(s, RIGHT_X, Emu(1250000), Emu(4200000), Emu(350000),
             [{"text": "HGINet 推理链（开源代码整理）", "size": 13, "color": PURPLE, "bold": True}])
    steps = [
        ("I¹ / I²", "双时相输入", GRAY),
        ("PVT-v2-B2", "共享骨干", RGBColor(0x5B, 0x2C, 0x6F)),
        ("MPA", "多层感知聚合", PURPLE),
        ("TC", "跨时相关增强", PURPLE),
        ("SDI", "差异图交互", PURPLE),
        ("三头输出", "Δ + Seg¹ + Seg²", DARK_RED),
    ]
    y = 1650000
    for i, (name, desc, fill) in enumerate(steps):
        add_rect(s, RIGHT_X, Emu(y), Emu(1500000), Emu(420000), fill)
        add_text(s, RIGHT_X, Emu(y), Emu(1500000), Emu(420000),
                 [{"text": name, "size": 12, "color": WHITE, "bold": True,
                   "font": FONT_EN if i < 2 else FONT}],
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(s, Emu(6200000), Emu(y), Emu(2600000), Emu(420000),
                 [{"text": desc, "size": 12, "color": INK}], valign=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            arr = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Emu(5200000), Emu(y + 410000), Emu(240000), Emu(160000))
            arr.fill.solid(); arr.fill.fore_color.rgb = RED; arr.line.fill.background()
        y += 580000
    add_text(s, L, Emu(5100000), FULL_W, Emu(1200000),
             [
                 [{"text": "与 HICD 的差异", "size": 13, "color": DARK_RED, "bold": True}],
                 [{"text": "HGINet 的「层级」在特征金字塔与图交互；HICD 的层级在标签空间（目标 ⊕ 状态）与实例可统计产品。二者都做语义变化，但标签组织不同。", "size": 12}],
                 [{"text": "阅读提示：写 related work 时区分「特征层级 / 标签层级 / 实例层级」三种层级语义。", "size": 12, "color": BLUE}],
             ])
    add_notes(s, "强调多任务三头输出，以及与HICD标签层级的差异；底部文字已下移避免与流程框重叠。")

    # ---------- 6 Architecture ----------
    s = blank_slide(prs)
    add_content_header(s, 6)
    add_red_label(s, "文章层次结构 / 方法总览：", top=Emu(820000))
    if flow and flow.exists():
        add_picture_fit(s, flow, L, Emu(1200000), FULL_W, Emu(3200000))
        add_text(s, L, Emu(4450000), FULL_W, Emu(250000),
                 [{"text": "图：官方开源仓库 flowchart（github.com/long123524/HGINet-torch）", "size": 11, "color": GRAY}])
    else:
        add_text(s, L, Emu(2000000), FULL_W, Emu(1200000),
                 [{"text": "（官方 flowchart 未能载入，见下页模块拆解）", "size": 16, "color": GRAY}],
                 align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(4800000), FULL_W, Emu(250000),
             [{"text": "五层解读", "size": 13, "color": RED, "bold": True}])
    notes = [
        ("输入层", "双时相影像"),
        ("特征层", "MPA 金字塔聚合（LAM+CBAM+门控融合）"),
        ("时序层", "TC 双向跨时相注意力，护未变化区"),
        ("变化层", "SDI 语义差分图卷积 + 通道门控"),
        ("输出层", "变化图 + 双时相语义分割"),
    ]
    y = 5100000
    for a, b in notes:
        add_text(s, L, Emu(y), FULL_W, Emu(280000),
                 [[{"text": f"{a}　", "size": 12, "color": PURPLE, "bold": True},
                   {"text": b, "size": 12, "color": INK}]])
        y += 280000
    add_notes(s, "用官方总览图串起输入到输出，并按五层口头展开。")

    # ---------- 7 Modules ----------
    s = blank_slide(prs)
    add_content_header(s, 7)
    add_red_label(s, "方法核心模块：", top=Emu(820000))
    modules = [
        (
            "MPA 多层感知聚合",
            "骨干 PVT-v2-B2 四级特征；每级 LAM（长短条纹注意）与 CBAM（通道/空间）并联；高层与低层细节门控融合（high + high⊙softmax(low)）。",
            "作用：多尺度可判别语义，服务分割头。",
            LIGHT_PURPLE, PURPLE,
        ),
        (
            "TC 时序相关模块",
            "双时相融合特征做双向 cross-attention：query 来自时相 A 粗先验池化，key/value 来自时相 B；残差系数 γ 可学习，初始为 0。",
            "作用：增强未变化区一致性，抑制伪变化。",
            LIGHT_BLUE, PURPLE,
        ),
        (
            "SDI 语义差分交互",
            "cz=|F²−F¹| 与 F¹、F² 拼接映射到 32 维；空间展平为图节点做 GCN；通道门控 σ(MLP(GAP(·))) 调制后卷积出变化图。",
            "作用：图学习刻画语义-差分交互，而非裸差分。",
            LIGHT_RED, DARK_RED,
        ),
    ]
    y = 1200000
    for title, body, role, fill, tcol in modules:
        add_rect(s, L, Emu(y), FULL_W, Emu(1350000), fill)
        add_text(s, Emu(380000), Emu(y + 60000), Emu(8200000), Emu(1250000),
                 [
                     [{"text": title, "size": 15, "color": tcol, "bold": True}],
                     [{"text": body, "size": 12, "color": INK}],
                     [{"text": role, "size": 12, "color": BLUE}],
                 ])
        y += 1450000
    add_text(s, L, Emu(5700000), FULL_W, Emu(600000),
             [
                 [{"text": "输出头：", "size": 13, "color": DARK_RED, "bold": True},
                  {"text": "out_feature / out_feature2 → 双时相语义；out_feature1 → 变化图；分割与变化联合训练。", "size": 12, "color": DARK_RED}],
                 [{"text": "实现核验：模块均见于官方 models.py（TC / SDI / LAM / CBAM / GCN）。", "size": 11, "color": GRAY}],
             ])
    add_notes(s, "模块页：MPA提特征、TC护未变化、SDI做差异图交互；均来自官方代码可核验实现。")

    # ---------- 8 Datasets ----------
    s = blank_slide(prs)
    add_content_header(s, 8)
    add_red_label(s, "常用数据集和评估方法：", top=Emu(820000))
    ds = [
        ("SECOND", "公开 SCD 基准\n城乡地物 from-to", "像元级语义变化"),
        ("HRSCD", "高分辨率语义变化\n土地利用/覆盖", "公开可复现"),
        ("Fuzhou", "合作方数据\n耕地→非耕地", "可外迁验证"),
        ("Xiamen", "域迁移验证\n同任务异城市", "Fuzhou 可迁入"),
    ]
    x = 280000
    cw = 2050000
    for name, desc, note in ds:
        add_rect(s, Emu(x), Emu(1250000), Emu(cw), Emu(1500000), LIGHT_GRAY)
        add_text(s, Emu(x + 80000), Emu(1320000), Emu(cw - 160000), Emu(1400000),
                 [
                     [{"text": name, "size": 14, "color": PURPLE, "bold": True, "font": FONT_EN}],
                     [{"text": desc, "size": 11}],
                     [{"text": note, "size": 11, "color": BLUE}],
                 ])
        x += 2150000
    add_text(s, L, Emu(2950000), Emu(4000000), Emu(300000),
             [{"text": "精度指标", "size": 13, "color": RED, "bold": True}])
    add_text(s, L, Emu(3280000), Emu(4000000), Emu(2000000),
             [
                 [{"text": "F1scd：", "size": 13, "color": BLUE, "bold": True, "font": FONT_EN},
                  {"text": "语义变化综合 F1", "size": 12}],
                 [{"text": "SeK：", "size": 13, "color": BLUE, "bold": True, "font": FONT_EN},
                  {"text": "分离变化/未变化的 Kappa", "size": 12}],
                 [{"text": "对比：", "size": 13, "color": BLUE, "bold": True},
                  {"text": "9 个既有方法；强调边界与抗伪变化", "size": 12}],
                 [{"text": "代码可复现 PVT-v2-B2 权重", "size": 12, "color": GRAY}],
             ])
    add_text(s, RIGHT_X, Emu(2950000), Emu(4200000), Emu(300000),
             [{"text": "与我们协议的对齐点", "size": 13, "color": RED, "bold": True}])
    add_text(s, RIGHT_X, Emu(3280000), Emu(4200000), Emu(2000000),
             [
                 [{"text": "HICD 已实现 SeK（bdd_protocol 等）", "size": 12}],
                 [{"text": "0617final 可报 tMIoU / sMIoU / SeK", "size": 12}],
                 [{"text": "SECOND 可作为公开基准对照", "size": 12}],
                 [{"text": "注意：HGINet 是类别 SCD，非 target-state 损毁层级", "size": 12, "color": DARK_RED}],
                 [{"text": "迁移结论可参照我们 Leave-one-city-out", "size": 12, "color": BLUE}],
             ])
    add_text(s, L, Emu(5600000), FULL_W, Emu(700000),
             [
                 [{"text": "阅读要点：", "size": 13, "color": DARK_RED, "bold": True},
                  {"text": "指标可写入 P1 主表协议说明；Fuzhou→Xiamen 迁移结论对 P2 跨城分析有参照价值。", "size": 12, "color": DARK_RED}],
             ])
    add_notes(s, "数据集四块+指标两块；强调SeK与SECOND对照对P1的意义。")

    # ---------- 9 Results ----------
    s = blank_slide(prs)
    add_content_header(s, 9)
    add_red_label(s, "实验对比结果（摘要报告最高值）：", top=Emu(820000))
    headers = ["数据集", "F1scd (%)", "SeK (%)", "设定要点"]
    rows = [
        ["SECOND", "59.48", "19.34", "公开 SCD，9 方法对比"],
        ["HRSCD", "64.12", "14.55", "高分辨率语义变化"],
        ["Fuzhou", "64.45", "18.28", "耕地→非耕地"],
        ["Xiamen", "84.93", "51.12", "同域较易 / 可迁移"],
    ]
    x0, y0 = 280000, 1250000
    col_w = [2000000, 1600000, 1600000, 3384000]
    x = x0
    for h, w in zip(headers, col_w):
        add_rect(s, Emu(x), Emu(y0), Emu(w), Emu(420000), PURPLE)
        align = PP_ALIGN.LEFT if h == "设定要点" else PP_ALIGN.CENTER
        add_text(s, Emu(x + (80000 if h == "设定要点" else 0)), Emu(y0), Emu(w), Emu(420000),
                 [{"text": h, "size": 13, "color": WHITE, "bold": True,
                   "font": FONT_EN if any(k in h for k in ("F1", "SeK", "SECOND")) else FONT}],
                 align=align, valign=MSO_ANCHOR.MIDDLE)
        x += w
    y = y0 + 420000
    for i, row in enumerate(rows):
        x = x0
        bg = LIGHT_PURPLE if i % 2 == 0 else WHITE
        for j, (val, w) in enumerate(zip(row, col_w)):
            add_rect(s, Emu(x), Emu(y), Emu(w), Emu(480000), bg, line=RGBColor(0xD0, 0xD0, 0xD0))
            color = DARK_RED if j in (1, 2) else INK
            align = PP_ALIGN.LEFT if j == 3 else PP_ALIGN.CENTER
            pad = 80000 if j == 3 else 0
            add_text(s, Emu(x + pad), Emu(y), Emu(w), Emu(480000),
                     [{"text": val, "size": 13, "color": color, "bold": j in (1, 2),
                       "font": FONT_EN if j < 3 else FONT}],
                     align=align, valign=MSO_ANCHOR.MIDDLE)
            x += w
        y += 480000
    # mini bar chart drawn with rects for F1scd
    add_text(s, L, Emu(3750000), Emu(4000000), Emu(280000),
             [{"text": "F1scd 对比示意（%）", "size": 12, "color": RED, "bold": True}])
    bar_data = [("SECOND", 59.48), ("HRSCD", 64.12), ("Fuzhou", 64.45), ("Xiamen", 84.93)]
    by = 4100000
    for name, val in bar_data:
        add_text(s, L, Emu(by), Emu(900000), Emu(280000),
                 [{"text": name, "size": 11, "font": FONT_EN}], valign=MSO_ANCHOR.MIDDLE)
        bw = int(val / 90.0 * 2400000)
        add_rect(s, Emu(1300000), Emu(by + 40000), Emu(bw), Emu(200000), PURPLE)
        add_text(s, Emu(1300000 + bw + 50000), Emu(by), Emu(700000), Emu(280000),
                 [{"text": f"{val:.2f}", "size": 11, "color": DARK_RED, "bold": True, "font": FONT_EN}],
                 valign=MSO_ANCHOR.MIDDLE)
        by += 320000
    add_text(s, RIGHT_X, Emu(3750000), Emu(4200000), Emu(2000000),
             [
                 [{"text": "定性结论（摘要）", "size": 13, "color": RED, "bold": True}],
                 [{"text": "• 变化区域与类别识别总体更优", "size": 12}],
                 [{"text": "• 缓解季节伪变化，边界更清晰", "size": 12}],
                 [{"text": "• Fuzhou→Xiamen 可迁移，跨城有鲁棒性", "size": 12}],
                 [{"text": "对我们的含义：", "size": 12, "color": DARK_RED, "bold": True}],
                 [{"text": "图交互+TC 可借鉴防伪变化；但 SECOND SeK 仅 19.34，SCD 仍难。", "size": 12, "color": DARK_RED}],
             ])
    add_text(s, L, Emu(6200000), FULL_W, Emu(300000),
             [{"text": "数据来源：Zotero 摘要所报四数据集最高 F1scd/SeK；开源代码用于结构核验，非复现实验表。", "size": 11, "color": GRAY}])
    add_notes(s, "报四数据集最高值，并指出SeK绝对值仍低，SCD任务本身难。")

    # ---------- 10 Comparison ----------
    s = blank_slide(prs)
    add_content_header(s, 10)
    add_red_label(s, "对照 HICD：同源「图 + 层级语义」，不同落点：", top=Emu(820000))
    add_rect(s, L, Emu(1250000), Emu(4050000), Emu(3400000), LIGHT_PURPLE)
    add_rect(s, RIGHT_X, Emu(1250000), Emu(4200000), Emu(3400000), LIGHT_BLUE)
    add_text(s, Emu(360000), Emu(1320000), Emu(3900000), Emu(3200000),
             [
                 [{"text": "HGINet（本文）", "size": 15, "color": PURPLE, "bold": True}],
                 [{"text": "图作用对象：特征", "size": 13, "bold": True}],
                 [{"text": "SDI-GCN 交互 F¹/F²/|F¹−F²| 空间节点", "size": 12}],
                 [{"text": "层级含义：特征金字塔", "size": 13, "bold": True}],
                 [{"text": "MPA 多尺度聚合，非标签本体", "size": 12}],
                 [{"text": "任务输出：类别 SCD", "size": 13, "bold": True}],
                 [{"text": "像元 from-to + 双时相分割", "size": 12}],
                 [{"text": "场景：城乡/耕地等公开 SCD", "size": 13, "bold": True}],
                 [{"text": "SECOND / HRSCD / 福州 / 厦门", "size": 12}],
                 [{"text": "贡献定位：特征交互增强变化判别", "size": 12, "color": BLUE}],
                 [{"text": "局限：无实例清单 / 标签拓扑弱", "size": 12, "color": DARK_RED}],
             ])
    add_text(s, Emu(4700000), Emu(1320000), Emu(4000000), Emu(3200000),
             [
                 [{"text": "HICD（本课题 v7）", "size": 15, "color": RGBColor(0x1F, 0x4E, 0x79), "bold": True}],
                 [{"text": "图作用对象：标签空间", "size": 13, "bold": True}],
                 [{"text": "GNNLabelStructure：15 目标 + 7 状态", "size": 12}],
                 [{"text": "层级含义：目标 ⊕ 损毁状态", "size": 13, "bold": True}],
                 [{"text": "条件解码 + 像素/实例协同", "size": 12}],
                 [{"text": "任务输出：战争损毁层级 SCD", "size": 13, "bold": True}],
                 [{"text": "变化场 + 实例 target–state", "size": 12}],
                 [{"text": "场景：0617final 多城战场", "size": 13, "bold": True}],
                 [{"text": "机场 / 港口 / 城乡异构", "size": 12}],
                 [{"text": "贡献定位：标签约束与可统计产品", "size": 12, "color": BLUE}],
                 [{"text": "缺口：通用 SCD 基线对比不足", "size": 12, "color": DARK_RED}],
             ])
    add_text(s, L, Emu(4800000), FULL_W, Emu(1500000),
             [
                 [{"text": "可写进论文的对照句", "size": 13, "color": RED, "bold": True}],
                 [{"text": "近期 SCD 工作（如 HGINet）开始用图结构建模双时相语义与差分特征的交互，改善复杂场景变化识别；但其「层级」主要停留在特征尺度，标签仍近似对等类别集合。面向战争损毁时，目标与状态强耦合且存在非法组合，需要把图结构引入标签与解码过程本身——这正是 HICD 的定位。", "size": 12}],
                 [{"text": "实验动作：", "size": 12, "color": DARK_RED, "bold": True},
                  {"text": "P1 基线可纳入 HGINet；消融增加「特征图交互 vs 标签图结构」对照说明。", "size": 12, "color": DARK_RED}],
             ])
    add_notes(s, "核心页：特征图GCN与标签图GNN的差异，直接服务related work与P1基线设计。")

    # ---------- 11 Challenges ----------
    s = blank_slide(prs)
    add_content_header(s, 11)
    add_red_label(s, "挑战与展望：", top=Emu(820000))
    points = [
        ("图建模粒度", "图节点是空间/通道展平特征，未显式建模类别/实例节点，对损毁状态细粒度表达有限。"),
        ("实例级缺位", "输出仍是像元变化图+分割，难以直接支撑「可清点、可制图」的战损产品。"),
        ("类别不平衡", "公开 SCD 相对均衡；战争损毁长尾（弹坑、塔台）会进一步压低 SeK。"),
        ("域迁移边界", "Fuzhou→Xiamen 说明同任务跨城可迁；跨传感器/跨冲突区需 Leave-one-out 分析。"),
        ("复现成本", "开源可复现，但标签映射与预处理对齐 SECOND 协议仍需工作量。"),
        ("正交机会", "特征图交互与标签图约束可叠加，支撑层级损毁 SCD 的扩展路径。"),
    ]
    y = 1250000
    for t, b in points:
        add_text(s, L, Emu(y), Emu(1400000), Emu(700000),
                 [{"text": t, "size": 13, "color": BLUE, "bold": True}], valign=MSO_ANCHOR.MIDDLE)
        add_text(s, Emu(1800000), Emu(y), Emu(7000000), Emu(700000),
                 [{"text": b, "size": 12, "color": INK}], valign=MSO_ANCHOR.MIDDLE)
        y += 780000
    add_text(s, L, Emu(6000000), FULL_W, Emu(400000),
             [{"text": "展望：把特征级图交互与标签级图约束、实例级解码分层组合。", "size": 12, "color": DARK_RED}])
    add_notes(s, "六点挑战均映射到我们论文框架的缺口。")

    # ---------- 12 Summary ----------
    s = blank_slide(prs)
    add_content_header(s, 12)
    add_red_label(s, "总结与对课题的行动项：", top=Emu(820000))
    add_text(s, L, Emu(1250000), FULL_W, Emu(350000),
             [{"text": "可借鉴三点", "size": 14, "color": RED, "bold": True}])
    takeaways = [
        "① 多任务联合（分割 + 变化）优于纯变化头——对应像素头/状态头联训。",
        "② 显式建模未变化区时序相关（TC）有助压伪变化——可与频域/NFA 互补消融。",
        "③ 差异特征上的图交互（SDI）提供 Transformer 之外的变化建模路线。",
    ]
    y = 1600000
    for t in takeaways:
        add_text(s, L, Emu(y), FULL_W, Emu(380000), [{"text": t, "size": 13, "color": INK}])
        y += 400000
    add_text(s, L, Emu(2900000), FULL_W, Emu(350000),
             [{"text": "落地行动（对齐 paper_order_tasks / baseline_compare_plan）", "size": 14, "color": RED, "bold": True}])
    actions = [
        ("P0", "将 HGINet 列入 P1 公开基线候选（优先 SECOND 协议），写清映射与不可比项。"),
        ("P0", "相关工作增加「特征图交互 vs 标签图结构」对照段，服务 C2/P1 创新点。"),
        ("P1", "消融增加时序相关/伪变化抑制组件对比（TC 类机制 vs 频域/时序模块）。"),
        ("P1", "跟踪同组 BGSNet（ISPRS 2025）与 MISGNet 等 multilevel SCD。"),
    ]
    y = 3300000
    for tag, body in actions:
        add_rect(s, L, Emu(y), Emu(750000), Emu(400000), DARK_RED if tag == "P0" else PURPLE)
        add_text(s, L, Emu(y), Emu(750000), Emu(400000),
                 [{"text": tag, "size": 12, "color": WHITE, "bold": True, "font": FONT_EN}],
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(s, Emu(1200000), Emu(y), Emu(7600000), Emu(400000),
                 [{"text": body, "size": 12, "color": INK}], valign=MSO_ANCHOR.MIDDLE)
        y += 500000
    add_rect(s, L, Emu(5450000), FULL_W, Emu(1000000), LIGHT_RED)
    add_text(s, Emu(380000), Emu(5550000), Emu(8300000), Emu(850000),
             [
                 [{"text": "精读结论", "size": 13, "color": DARK_RED, "bold": True}],
                 [{"text": "HGINet 是 ISPRS 上「层级语义 + 图交互」SCD 的强基线与写作参照，但其层级在特征而非标签/实例；HICD 应强调标签拓扑、目标–状态耦合与区域制图出口。", "size": 12, "color": DARK_RED}],
             ])
    add_notes(s, "收束到可执行P0/P1行动，便于周报跟进。")

    prs.save(str(OUT))
    print("saved", OUT)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    build()
