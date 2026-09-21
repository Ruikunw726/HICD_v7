# -*- coding: utf-8 -*-
"""Rebuild HGINet reading PPT using full PDF text (Zotero QBLZQPBN)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(r"F:\mambacd\home\HICD_v7\docs\ISPRS2024_精读_王瑞坤.pptx")
ASSETS = Path(r"F:\mambacd\home\HICD_v7\docs\assets")
FIG_FLOW = ASSETS / "hginet_fig_workflow.png"
FIG_MPANET = ASSETS / "hginet_fig_mpanet.png"
FIG_TC = ASSETS / "hginet_fig_tc_sdi.png"
FIG_FLOW2 = ASSETS / "hginet_flowchart.png"
if not FIG_FLOW.exists() and FIG_FLOW2.exists():
    FIG_FLOW = FIG_FLOW2

PURPLE = RGBColor(0x6E, 0x0F, 0x6D)
RED = RGBColor(0xFF, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x00, 0x00, 0xFF)
DARK_RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_PURPLE = RGBColor(0xF3, 0xE8, 0xF3)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xF8)
LIGHT_RED = RGBColor(0xFB, 0xEB, 0xEB)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)

FONT = "微软雅黑"
FONT_EN = "Times New Roman"
L = Emu(280000)
R = Emu(8864000)
FULL_W = Emu(R - L)
HALF_W = Emu(4050000)
RIGHT_X = Emu(4600000)


def set_cjk_font(run, font_name=FONT):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    successors = {
        "a:ea": ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
        "a:cs": ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
    }
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, *successors[tag])
        el.set("typeface", font_name)


def style_run(run, size=16, bold=False, color=INK, font_name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_cjk_font(run, font_name)


def add_text(slide, left, top, width, height, segments, *, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, word_wrap=True, fill=None, line=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = valign
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        if line is None:
            box.line.fill.background()
    if line is not None:
        box.line.color.rgb = line
    if isinstance(segments, str):
        segments = [(segments, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    for text, kw in segments:
        if text == "\n":
            p = tf.add_paragraph()
            p.alignment = kw.get("align", align)
            if "level" in kw:
                p.level = kw["level"]
            continue
        run = p.add_run()
        run.text = text
        style_run(run, size=kw.get("size", 16), bold=kw.get("bold", False),
                  color=kw.get("color", INK), font_name=kw.get("font", FONT))
    return box


def add_bullets(slide, left, top, width, height, items, size=15, color=INK, fill=None):
    """items: list of (text, opts) or str"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.fill.background()
    first = True
    for item in items:
        if isinstance(item, str):
            text, opts = item, {}
        else:
            text, opts = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.level = opts.get("level", 0)
        # bullet prefix
        prefix = "• " if p.level == 0 else "– "
        run = p.add_run()
        run.text = prefix + text
        style_run(run, size=opts.get("size", size), bold=opts.get("bold", False),
                  color=opts.get("color", color), font_name=opts.get("font", FONT))
    return box


def add_header(slide, page_no=None):
    add_text(slide, L, Emu(120000), Emu(2000000), Emu(400000),
             [("论文阅读", {"size": 22, "bold": True, "color": PURPLE})])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, Emu(520000), FULL_W, Emu(20000))
    line.fill.solid()
    line.fill.fore_color.rgb = PURPLE
    line.line.fill.background()
    if page_no is not None:
        add_text(slide, Emu(8200000), Emu(6900000), Emu(600000), Emu(280000),
                 [(str(page_no), {"size": 11, "color": GRAY})], align=PP_ALIGN.RIGHT)


def add_label_block(slide, top, label, body_segments, *, height=Emu(1400000), fill=None):
    add_text(slide, L, top, FULL_W, Emu(280000),
             [(label, {"size": 14, "bold": True, "color": RED})])
    add_text(slide, L, Emu(top + 280000), FULL_W, Emu(height - 280000),
             body_segments, fill=fill)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_picture_fit(slide, path, left, top, width, height):
    if not Path(path).exists():
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 1 cover
    s = new_slide(prs)
    add_text(s, L, Emu(1500000), FULL_W, Emu(800000),
             [("论文精读 · 组会汇报", {"size": 18, "color": PURPLE, "bold": True})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(2200000), FULL_W, Emu(1200000),
             [("Semantic change detection using a\nhierarchical semantic graph interaction network\nfrom high-resolution remote sensing images",
               {"size": 22, "bold": True, "color": INK, "font": FONT_EN})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(3800000), FULL_W, Emu(1000000),
             [("HGINet · ISPRS Journal of Photogrammetry and Remote Sensing 211 (2024) 318–335",
               {"size": 14, "color": GRAY}),
              ("\n", {}),
              ("Jiang Long, Mengmeng Li*, Xiaoqin Wang, Alfred Stein  |  Fuzhou University / ITC",
               {"size": 13, "color": GRAY}),
              ("\n", {}),
              ("DOI: 10.1016/j.isprsjprs.2024.04.012  |  Zotero key: WY3KTIRY / PDF: QBLZQPBN",
               {"size": 12, "color": GRAY}),
              ("\n", {}),
              ("汇报人：王瑞坤    数据源：Zotero 本地 PDF 全文（18 页）",
               {"size": 13, "color": INK})],
             align=PP_ALIGN.CENTER)
    notes(s, "精读对象为 Long et al. 2024 ISPRS J 的 HGINet。本次基于 Zotero 附件 PDF 全文，而非仅摘要。")

    # 2 why this paper
    s = new_slide(prs)
    add_header(s, 2)
    add_text(s, L, Emu(650000), FULL_W, Emu(400000),
             [("为什么精读这篇（给马磊/组会）", {"size": 24, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1100000), FULL_W, Emu(2200000), [
        ("顶刊 ISPRS J. 2024，作者含 Alfred Stein，方法开源：github.com/long123524/HGINet-torch", {}),
        ("问题直接打在 SCD 痛点上：双时相语义特征与差分特征的时序相关（TC）建模不足", {}),
        ("核心动作是「层级 + 图交互」：特征层之间用 GCN 建模，而非只做像素差", {}),
        ("与我们 HICD 最同构的一点：都用图结构处理「不规则/多层」关系——但图建在特征上，不是标签空间上", {"color": DARK_RED, "bold": True}),
        ("任务仍是城市地表 from–to 语义变化，不是战争 target×state；对照后能讲清我们的差异点", {}),
    ], size=15)
    add_label_block(s, Emu(3600000), "一句话问题：",
                    [("高分辨率影像上，如何同时更准地分清「哪里变了」和「变成什么类别」？尤其是复杂城市场景里，未变化区域容易被误判。",
                      {"size": 15, "bold": True})],
                    height=Emu(1000000), fill=LIGHT_PURPLE)
    add_label_block(s, Emu(4800000), "和我们当前工作的接口：",
                    [("像素级 SCD 基线可对照 SECOND 协议；图建模思路可支撑「标签/组合约束」叙事；TC 强化未变化区，对应我们 state=no-change 主导的问题。",
                      {"size": 14})],
                    height=Emu(1200000), fill=LIGHT_BLUE)
    notes(s, "开场对齐马磊要求：选文要有知识增量，并能落到自己研究。")

    # 3 citation card
    s = new_slide(prs)
    add_header(s, 3)
    add_text(s, L, Emu(650000), FULL_W, Emu(400000),
             [("文献卡片与选文逻辑", {"size": 24, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1100000), HALF_W, Emu(4500000), [
        ("题目：Hierarchical semantic graph interaction network (HGINet)", {"bold": True}),
        ("期刊：ISPRS J. Photogramm. Remote Sens. 211 (2024) 318–335", {}),
        ("单位：福州大学 数字中国研究院（教育部空间数据挖掘重点实验室）+ 特文特大学 ITC", {}),
        ("关键词：SCD / Temporal correlations / Semantic difference interaction", {}),
        ("代码：long123524/HGINet-torch", {}),
        ("库内：collection「变化检测」，itemKey=WY3KTIRY", {}),
        ("PDF：C:\\Users\\Ruiku\\Zotero\\storage\\QBLZQPBN\\…pdf（21.9 MB，18 页）", {"size": 13, "color": GRAY}),
    ], size=14)
    add_bullets(s, RIGHT_X, Emu(1100000), HALF_W, Emu(4500000), [
        ("官方四条贡献（全文 §1）", {"bold": True, "color": RED}),
        ("1) 面向多场景的稳健 SCD 方法", {}),
        ("2) 轻量 SDI：用图学习刻画多层特征交互", {}),
        ("3) TC 模块：交叉学习增强未变化区", {}),
        ("4) MPANet：金字塔式多层级感知聚合特征提取", {}),
        ("", {}),
        ("对比基线 9 个", {"bold": True, "color": RED}),
        ("BIT / GAS-Net / HRSCD-3 / HRSCD-4 / FCCDN / SSCDl / Bi-SRNet / PCFN / MTSCD-Net", {"size": 13}),
    ], size=14)
    notes(s, "强调本次有全文支撑，不是二手摘要。")

    # 4 problem
    s = new_slide(prs)
    add_header(s, 4)
    add_text(s, L, Emu(650000), FULL_W, Emu(400000),
             [("研究背景：SCD 为什么难（全文动机）", {"size": 24, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1100000), FULL_W, Emu(2800000), [
        ("BCD 只回答变/不变；SCD 还要给出 from–to 类别变化", {}),
        ("单分支：把变化类型当独立类 → 类别数≈语义类平方，严重不平衡（全文明确点出）", {}),
        ("双分支后分类比较：依赖两次分割精度，边界误差累积", {}),
        ("多任务（分割+变化）更合适，但多数工作忽略双时相特征的时序相关 TC", {"bold": True, "color": DARK_RED}),
        ("CNN 感受野有限：长程上下文弱、多层特征全局交互弱 → 复杂场景误检", {}),
        ("未变化区：双时相特征应高相似；现有多任务网络未显式利用 → 假变化/漏检", {}),
    ], size=15)
    add_label_block(s, Emu(4200000), "对我们战争场景的映射：",
                    [("0617final 不是 from–to，而是 target×state；未变化背景占比极高（~70%+）。HGINet 强调「先稳住未变化」与「多层交互」，对我们 state=no-change 与长尾变化类都有参考价值。",
                      {"size": 14})],
                    height=Emu(1400000), fill=LIGHT_RED)
    notes(s, "动机段落来自论文 Introduction，已对照全文。")

    # 5 task & data
    s = new_slide(prs)
    add_header(s, 5)
    add_text(s, L, Emu(650000), FULL_W, Emu(400000),
             [("任务定义与数据（全文 §3）", {"size": 24, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1100000), HALF_W, Emu(4000000), [
        ("任务：多任务 SCD = 双时相语义分割 + 变化检测", {"bold": True}),
        ("输入：双时相高分影像对", {}),
        ("输出：SS1 / SS2 分割图 + 变化图（from–to）", {}),
        ("损失：Lt = 0.5(Ls1+Ls2) + Lc + Lsc", {"color": BLUE, "bold": True}),
        ("Lsc：余弦损失约束未变化区一致性（yc=0 拉近 / yc=1 推开）", {}),
        ("训练：RTX 3090，SGD，batch=8，80 epochs，weight decay 5e-4", {"size": 13, "color": GRAY}),
        ("数据增强：随机旋转、翻转", {"size": 13, "color": GRAY}),
    ], size=14)
    add_bullets(s, RIGHT_X, Emu(1100000), HALF_W, Emu(4000000), [
        ("四个数据集", {"bold": True, "color": RED}),
        ("SECOND：公开 SCD 基准（变化类型多、复杂）", {}),
        ("HRSCD：公开/城市语义变化", {}),
        ("Fuzhou：Google Earth 0.5 m，仓山/闽侯等；训练 6104 / 测试 2615（256²）", {}),
        ("Xiamen：同安/翔安；训练 4542 / 测试 1946；7:3 划分", {}),
        ("Fuzhou/Xiamen 七类：unchanged, tree, cropland, building, water, bare land, road", {"size": 13}),
        ("指标：mIoU、SeK、F1scd、OS", {"bold": True}),
        ("作者提醒：SECOND/HRSCD 存在较多错误标注（Discussion & Fig.15）", {"color": DARK_RED}),
    ], size=14)
    notes(s, "数据细节来自全文 §3，非官方 repo 猜测。")

    # 6 method overview + figure
    s = new_slide(prs)
    add_header(s, 6)
    add_text(s, L, Emu(600000), FULL_W, Emu(350000),
             [("方法总览：HGINet 三大组件（全文 Fig.1）", {"size": 22, "bold": True, "color": PURPLE})])
    if FIG_FLOW.exists():
        add_picture_fit(s, FIG_FLOW, L, Emu(1000000), FULL_W, Emu(3200000))
    add_bullets(s, L, Emu(4400000), FULL_W, Emu(2200000), [
        ("① MPANet（Siamese）：PVT-v2 骨干 + 多层级感知聚合 → 双时相语义特征", {}),
        ("② TC 模块：非局部交叉学习建模双时相时序相关 → 强化未变化区识别", {}),
        ("③ SDI 模块：GCN 建模「语义特征层 × 差分特征层 × 二者组合」交互 → 变化区/类", {"color": DARK_RED}),
        ("多任务约束：用变化定位约束语义类别学习，而不是两套头各学各的", {}),
    ], size=14)
    notes(s, "Fig.1 取自论文 PDF 第2页渲染图。")

    # 7 modules detail
    s = new_slide(prs)
    add_header(s, 7)
    add_text(s, L, Emu(600000), FULL_W, Emu(350000),
             [("模块拆解（全文 §2.1–2.3）", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1000000), HALF_W, Emu(5200000), [
        ("MPANet 特征提取", {"bold": True, "color": RED}),
        ("骨干：ImageNet 预训练 PVT-V2（空间缩减注意力，金字塔）", {}),
        ("F1 细粒度；F2–F4 高层语义", {}),
        ("MSCF：LAM（横/纵/双对角条带卷积）+ CBAM 近程注意力", {"color": BLUE}),
        ("SFSR：用 F1 的 Softmax 注意力 refine 高层特征，补边界细节", {"color": BLUE}),
        ("TC 时序相关", {"bold": True, "color": RED}),
        ("双时相特征投影为 Q/K/V（通道下采样 r=2）", {}),
        ("非局部网络全局相关；交叉学习迁移知识 → 未变化区更稳", {}),
    ], size=13)
    add_bullets(s, RIGHT_X, Emu(1000000), HALF_W, Emu(5200000), [
        ("SDI 语义差分交互（核心）", {"bold": True, "color": RED}),
        ("图节点 ≈ 不同特征层表示（语义 / 差分 / 联合）", {}),
        ("GCN 量化层间交互，全局视角建模复杂场景变化", {"color": BLUE}),
        ("消融显示 SDI 贡献最大：SECOND 上 OS +3.91%，F1scd +4.57%", {"bold": True, "color": DARK_RED}),
        ("模块轻量：thop 统计参数量/计算量增量小", {}),
        ("", {}),
        ("与 HICD 的「图」差在哪", {"bold": True, "color": RED}),
        ("HGINet：图建在特征层交互", {}),
        ("HICD：图建在 target×state 标签合法性", {"color": DARK_RED}),
        ("可讲成「特征图 GCN vs 标签图 GNN」", {"bold": True}),
    ], size=13)
    notes(s, "模块细节来自 §2 正文公式与消融，不是 repo 目录名堆砌。")

    # 8 results
    s = new_slide(prs)
    add_header(s, 8)
    add_text(s, L, Emu(600000), FULL_W, Emu(350000),
             [("主要结果（全文摘要 + §4 + §6）", {"size": 22, "bold": True, "color": PURPLE})])
    # table-like blocks
    add_text(s, L, Emu(1050000), FULL_W, Emu(300000),
             [("四数据集最优（相对 9 个基线）", {"size": 14, "bold": True, "color": RED})])
    # simple table via shapes
    headers = ["数据集", "F1scd", "SeK"]
    rows = [
        ["SECOND", "59.48%", "19.34%"],
        ["HRSCD", "64.12%", "14.55%"],
        ["Fuzhou", "64.45%", "18.28%"],
        ["Xiamen", "84.93%", "51.12%"],
    ]
    col_w = [Emu(2200000), Emu(2000000), Emu(2000000)]
    x0 = L
    y0 = Emu(1400000)
    for j, h in enumerate(headers):
        add_text(s, x0, y0, col_w[j], Emu(320000),
                 [(h, {"size": 13, "bold": True, "color": WHITE})],
                 fill=RGBColor(0x6E, 0x0F, 0x6D), align=PP_ALIGN.CENTER)
        x0 = Emu(x0 + col_w[j])
    for i, row in enumerate(rows):
        x0 = L
        y = Emu(y0 + 320000 * (i + 1))
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        for j, cell in enumerate(row):
            add_text(s, x0, y, col_w[j], Emu(320000),
                     [(cell, {"size": 13, "bold": j > 0})],
                     fill=bg, align=PP_ALIGN.CENTER)
            x0 = Emu(x0 + col_w[j])
    add_bullets(s, L, Emu(3200000), FULL_W, Emu(3000000), [
        ("相对 9 个 SOTA：F1scd 提升 1.28–11.53 / 1.78–30.21 / 10.32–25.42 / 1.34–17.69（四数据集，§6）", {}),
        ("消融（SECOND）：去掉全部增强模块，OS −4.91%、F1scd −6.06%；SDI 单独贡献最大", {"color": DARK_RED}),
        ("骨干对比：MPANet > 纯 PVT-V2 > Efficient-B3 > ResNet34（SeK/F1scd 最高）", {}),
        ("Fuzhou→Xiamen 迁移成功，作者称对季节/光谱差异适应性好", {}),
        ("变化类型少、目标更规整时精度更高（Xiamen）；小而乱的目标更难（Fuzhou）", {}),
        ("局限：边界模糊、复杂背景包围的目标仍难；公开数据标签噪声大", {"color": DARK_RED}),
    ], size=14)
    notes(s, "数字来自摘要与结论，已与全文交叉核对。")

    # 9 compare HICD
    s = new_slide(prs)
    add_header(s, 9)
    add_text(s, L, Emu(600000), FULL_W, Emu(350000),
             [("与我们工作的对照（基于全文，不空谈）", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1000000), HALF_W, Emu(5200000), [
        ("相同点", {"bold": True, "color": RED}),
        ("双时相高分影像变化理解", {}),
        ("多任务/多头：类别 + 变化信息联合", {}),
        ("关注未变化主导 + 变化稀疏", {}),
        ("都用 SeK 等分离变化评估", {}),
        ("都承认复杂场景下边界与细小目标难", {}),
        ("图/结构化建模「关系」", {}),
    ], size=14)
    add_bullets(s, RIGHT_X, Emu(1000000), HALF_W, Emu(5200000), [
        ("不同点（对组会更关键）", {"bold": True, "color": RED}),
        ("语义空间：from–to 地类 vs target×state 战损", {"color": DARK_RED}),
        ("图的位置：特征层 GCN vs 标签合法组合 GNN", {"color": DARK_RED}),
        ("目标体系：通用 6–7 类 vs 15 类特殊目标长尾", {}),
        ("应用出口：城市扩张/农地 vs 战损制图与清点", {}),
        ("可借鉴：TC 强化 no-change；SDI 的层间交互思想", {"color": BLUE}),
        ("可借鉴：LAM 条带卷积对跑道/道路等线状目标", {"color": BLUE}),
        ("不可直接搬：SECOND 结论不能外推到弹坑/掩体", {"bold": True}),
    ], size=14)
    notes(s, "对照要服务马磊：讲清差异，避免被问「你和这篇有啥不一样」。")

    # 10 discussion for 马磊
    s = new_slide(prs)
    add_header(s, 10)
    add_text(s, L, Emu(600000), FULL_W, Emu(350000),
             [("给马磊的讨论点 & 对我们的行动项", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1000000), FULL_W, Emu(2800000), [
        ("讨论 1：图建在「特征层」还是「标签/目标–状态关系」更划算？战争组合稀疏时，标签图是否比特征 GCN 更对症？", {"bold": True}),
        ("讨论 2：识别尺度——弹坑/战壕是单点检测，还是「成群区域」作为一层语义？HGINet 的对象级 SCD 粒度对不上我们的实例清单需求。", {"bold": True}),
        ("讨论 3：他们承认 SECOND/HRSCD 标签噪声；我们战损标注质量参差，是否更需要样本表征与噪声稳健，而不是再堆结构？", {}),
    ], size=14)
    add_bullets(s, L, Emu(4000000), FULL_W, Emu(2400000), [
        ("行动项（组会后）", {"bold": True, "color": RED}),
        ("P1 对照实验：在 SECOND 协议下把 HGINet 作为像素级 SCD 基线之一（有官方代码）", {}),
        ("像素分支：评估 TC/交叉学习类「未变化一致性」损失，对 state=no-change 是否有增益", {}),
        ("线状目标（跑道/滑行道）可试 LAM/条带注意力，而非先上重型混合 backbone", {}),
        ("组会口径：先讲清 from–to vs target×state，再讲图的用法差异", {"color": DARK_RED}),
    ], size=14)
    notes(s, "对齐马磊：问题讲明白 + 落到下一步工作。")

    # 11 summary
    s = new_slide(prs)
    add_header(s, 11)
    add_text(s, L, Emu(600000), FULL_W, Emu(350000),
             [("精读总结", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1100000), FULL_W, Emu(3500000), [
        ("HGINet 把 SCD 拆成：强特征（MPANet）+ 稳未变（TC）+ 图交互差分（SDI）", {"bold": True}),
        ("全文证据：SDI 消融最敏感；Transformer 特征提取优于纯 CNN；跨城迁移可用", {}),
        ("对我们：这是「城市 from–to SCD」的强基线与思路参考，不是战损问题的直接答案", {"color": DARK_RED}),
        ("最大启发不是再抄一个 GCN，而是：关系建模放哪里（特征 / 标签 / 实例）决定方法是否对症", {"bold": True}),
        ("下一步：SECOND 基线复现 + 把 TC 式未变化约束接到像素级；标签图 GNN 继续作为我们的主贡献叙事", {}),
    ], size=15)
    add_label_block(s, Emu(4800000), "参考文献：",
                    [("Long J., Li M., Wang X., Stein A., 2024. ISPRS J. Photogramm. Remote Sens. 211, 318–335. doi:10.1016/j.isprsjprs.2024.04.012",
                      {"size": 13, "color": GRAY, "font": FONT_EN})],
                    height=Emu(800000), fill=LIGHT_GRAY)
    notes(s, "总结页控制在三条可记住的信息。")

    prs.save(str(OUT))
    print("SAVED", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
