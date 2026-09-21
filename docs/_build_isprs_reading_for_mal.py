# -*- coding: utf-8 -*-
"""Group-meeting paper-reading PPT for HGINet (ISPRS 2024) — for 马磊."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(r"F:\mambacd\home\HICD_v7\docs\ISPRS2024_精读_王瑞坤.pptx")
ASSETS = Path(r"F:\mambacd\home\HICD_v7\docs\assets")
FIG_ARCH = ASSETS / "hginet_flowchart.png"          # Fig.1 真实架构图（宽）
FIG_ARCH_S = ASSETS / "hginet_flowchart_small.png"
FIG_P3 = ASSETS / "fig_p3_full_top.png"
FIG_P4 = ASSETS / "fig_p4_full_mid.png"
FIG_P5 = ASSETS / "p05_mid.png"

PURPLE = RGBColor(0x6E, 0x0F, 0x6D)
RED = RGBColor(0xFF, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x00, 0x00, 0xFF)
DARK_RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_PURPLE = RGBColor(0xF3, 0xE8, 0xF3)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xF8)
LIGHT_RED = RGBColor(0xFB, 0xEB, 0xEB)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
LIGHT_YELLOW = RGBColor(0xFF, 0xF8, 0xE1)

FONT = "微软雅黑"
FONT_EN = "Times New Roman"
L = Emu(280000)
R = Emu(8864000)
FULL_W = Emu(R - L)
HALF_W = Emu(4100000)
RIGHT_X = Emu(4700000)


def set_cjk_font(run, font_name=FONT):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    successors = {
        "a:ea": ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
        "a:cs": ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
    }
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, *successors[tag])
        el.set("typeface", font_name)


def style_run(run, size=16, bold=False, color=INK, font_name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_cjk_font(run, font_name)


def add_text(slide, left, top, width, height, segments, *, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, fill=None, line=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        if line is None:
            box.line.fill.background()
    if line is not None:
        box.line.color.rgb = line
    if isinstance(segments, str):
        segments = [(segments, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    for text, kw in segments:
        if text == "\n":
            p = tf.add_paragraph()
            p.alignment = kw.get("align", align)
            continue
        run = p.add_run()
        run.text = text
        style_run(run, size=kw.get("size", 16), bold=kw.get("bold", False),
                  color=kw.get("color", INK), font_name=kw.get("font", FONT))
    return box


def add_bullets(slide, left, top, width, height, items, size=15, color=INK, fill=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.fill.background()
    first = True
    for item in items:
        text, opts = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.level = opts.get("level", 0)
        prefix = "• " if p.level == 0 else "– "
        run = p.add_run()
        run.text = prefix + text
        style_run(run, size=opts.get("size", size), bold=opts.get("bold", False),
                  color=opts.get("color", color), font_name=opts.get("font", FONT))
    return box


def add_header(slide, page_no=None, title=None):
    add_text(slide, L, Emu(100000), Emu(2200000), Emu(360000),
             [("论文阅读", {"size": 20, "bold": True, "color": PURPLE})])
    if title:
        add_text(slide, Emu(2200000), Emu(120000), Emu(6200000), Emu(320000),
                 [(title, {"size": 12, "color": GRAY})], align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, Emu(480000), FULL_W, Emu(18000))
    line.fill.solid()
    line.fill.fore_color.rgb = PURPLE
    line.line.fill.background()
    if page_no is not None:
        add_text(slide, Emu(8200000), Emu(6950000), Emu(600000), Emu(250000),
                 [(str(page_no), {"size": 11, "color": GRAY})], align=PP_ALIGN.RIGHT)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def pic(slide, path, left, top, width, height=None):
    path = Path(path)
    if not path.exists():
        add_text(slide, left, top, width, Emu(800000),
                 [(f"[图缺失: {path.name}]", {"size": 12, "color": RED})])
        return None
    if height is None:
        return slide.shapes.add_picture(str(path), left, top, width=width)
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ========== 1 封面 ==========
    s = new_slide(prs)
    add_text(s, L, Emu(800000), FULL_W, Emu(400000),
             [("文献精读汇报", {"size": 16, "bold": True, "color": PURPLE})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(1300000), FULL_W, Emu(1400000),
             [("Semantic Change Detection using a\nHierarchical Semantic Graph Interaction Network",
               {"size": 24, "bold": True, "font": FONT_EN, "color": INK})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(2800000), FULL_W, Emu(400000),
             [("—— 高分辨率遥感影像语义变化检测中的层级语义图交互网络（HGINet）",
               {"size": 16, "color": INK})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(3400000), FULL_W, Emu(1400000),
             [("Jiang Long, Mengmeng Li, Xiaoqin Wang, Alfred Stein",
               {"size": 14, "font": FONT_EN, "color": GRAY}),
              ("\n", {}),
              ("ISPRS Journal of Photogrammetry and Remote Sensing, 2024, 211: 318–335",
               {"size": 14, "font": FONT_EN, "color": GRAY}),
              ("\n", {}),
              ("福州大学 · 特文特大学 ITC    DOI: 10.1016/j.isprsjprs.2024.04.012",
               {"size": 13, "color": GRAY}),
              ("\n", {}),
              ("代码开源：github.com/long123524/HGINet-torch", {"size": 13, "color": GRAY})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(5200000), FULL_W, Emu(600000),
             [("汇报人：王瑞坤", {"size": 16, "bold": True})],
             align=PP_ALIGN.CENTER)
    add_text(s, L, Emu(5700000), FULL_W, Emu(400000),
             [("组会 · 论文精读（基于 PDF 全文）", {"size": 12, "color": GRAY})],
             align=PP_ALIGN.CENTER)
    notes(s, "开场30秒：这篇是2024年ISPRS Journal上的语义变化检测工作，核心是用图交互建模双时相特征关系。")

    # ========== 2 目录 ==========
    s = new_slide(prs)
    add_header(s, 2)
    add_text(s, L, Emu(650000), FULL_W, Emu(400000),
             [("汇报提纲", {"size": 26, "bold": True, "color": PURPLE})])
    items = [
        "一、研究背景与要解决的问题",
        "二、方法：HGINet 怎么做",
        "三、实验设置与主要结果",
        "四、消融分析与作者讨论",
        "五、这篇工作好在哪里 / 局限在哪",
        "六、对变化检测研究的几点启发",
    ]
    y = Emu(1300000)
    for i, t in enumerate(items):
        add_text(s, Emu(800000), y, Emu(8000000), Emu(500000),
                 [(t, {"size": 20, "bold": i < 3, "color": INK if i < 3 else GRAY})],
                 fill=LIGHT_PURPLE if i < 3 else None)
        y = Emu(y + 650000)
    notes(s, "提纲按论文逻辑走：问题—方法—实验—讨论—启发。")

    # ========== 3 背景与问题 ==========
    s = new_slide(prs)
    add_header(s, 3, "HGINet · ISPRS 2024")
    add_text(s, L, Emu(600000), FULL_W, Emu(400000),
             [("一、研究背景：语义变化检测难在哪", {"size": 22, "bold": True, "color": PURPLE})])
    add_text(s, L, Emu(1050000), FULL_W, Emu(280000),
             [("任务层次", {"size": 14, "bold": True, "color": RED})])
    add_bullets(s, L, Emu(1350000), HALF_W, Emu(2200000), [
        ("二值变化检测（BCD）：只判断「变 / 不变」", {}),
        ("语义变化检测（SCD）：还要回答「从什么变成什么」（from–to）", {"bold": True}),
        ("SCD 信息量更大，但要同时做对「位置」和「类别」", {"color": DARK_RED}),
    ], size=14)
    add_text(s, RIGHT_X, Emu(1050000), HALF_W, Emu(280000),
             [("已有方法的主要问题（论文动机）", {"size": 14, "bold": True, "color": RED})])
    add_bullets(s, RIGHT_X, Emu(1350000), HALF_W, Emu(2200000), [
        ("单分支：变化类型当独立类 → 类别爆炸、严重不平衡", {}),
        ("双分支：先分类再比较，边界误差会累积", {}),
        ("多任务较合理，但多数工作忽略双时相特征之间的时序相关", {"bold": True}),
        ("CNN 长程上下文弱，复杂城区易误判未变化区域", {}),
    ], size=14)
    add_text(s, L, Emu(3700000), FULL_W, Emu(2800000),
             [("一句话问题：", {"size": 15, "bold": True, "color": DARK_RED}),
              ("在高分辨率影像上，如何把「哪里变了」和「变成什么类」同时做准，并显式利用双时相特征之间的相关关系，减少复杂场景中的假变化？",
               {"size": 15, "bold": True})],
             fill=LIGHT_YELLOW)
    add_text(s, L, Emu(4800000), FULL_W, Emu(1200000),
             [("作者贡献（论文 §1 原文归纳）：\n① 稳健的多场景 SCD 方法　② 轻量 SDI（图交互）　③ TC 时序相关模块　④ MPANet 多层级特征提取",
               {"size": 13, "color": INK})],
             fill=LIGHT_GRAY)
    notes(s, "先讲清SCD相对BCD多了一层难度，再点出论文抓住的缺口：时序相关与多层特征交互。")

    # ========== 4 方法总览 + 真图 ==========
    s = new_slide(prs)
    add_header(s, 4, "HGINet · 方法")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("二、方法总览：HGINet 整体框架", {"size": 22, "bold": True, "color": PURPLE})])
    # 真实架构图 Fig.1
    pic(s, FIG_ARCH, L, Emu(1000000), FULL_W, Emu(2800000))
    add_text(s, L, Emu(3900000), FULL_W, Emu(280000),
             [("图：论文/官方代码中的 HGINet 架构（MPANet → TC → SDI）", {"size": 11, "color": GRAY})])
    add_bullets(s, L, Emu(4300000), FULL_W, Emu(2200000), [
        ("多任务网络：语义分割 + 变化检测联合学习，用变化定位约束语义类别", {}),
        ("① MPANet（Siamese 共享权重）：从 T1/T2 提取多层级语义特征", {}),
        ("② TC 模块：建模双时相特征的时序相关，强化「未变化区」识别", {"color": DARK_RED}),
        ("③ SDI 模块：用 GCN 建模语义特征、差分特征及其组合的层间交互", {"color": DARK_RED}),
        ("输出：T1/T2 语义分割图 + 语义变化图（from–to）", {}),
    ], size=14)
    notes(s, "指着架构图从左到右讲：输入双时相→特征提取→TC→SDI→输出分割与变化。这页是全场重点。")

    # ========== 5 模块细节 ==========
    s = new_slide(prs)
    add_header(s, 5, "HGINet · 方法细节")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("二、三个关键模块在做什么", {"size": 22, "bold": True, "color": PURPLE})])
    # left column MPANet
    add_text(s, L, Emu(1000000), HALF_W, Emu(300000),
             [("1. MPANet 特征提取", {"size": 15, "bold": True, "color": RED})])
    add_bullets(s, L, Emu(1350000), HALF_W, Emu(2400000), [
        ("骨干：ImageNet 预训练 PVT-V2（金字塔视觉 Transformer）", {}),
        ("多层级特征：F1 偏细节，F2–F4 偏语义", {}),
        ("MSCF：远程条带注意力 LAM + 近程 CBAM，融合上下文", {}),
        ("SFSR：用低层细节 refine 高层语义，改善边界", {}),
    ], size=13)
    add_text(s, RIGHT_X, Emu(1000000), HALF_W, Emu(300000),
             [("2. TC 时序相关模块", {"size": 15, "bold": True, "color": RED})])
    add_bullets(s, RIGHT_X, Emu(1350000), HALF_W, Emu(2400000), [
        ("双时相特征映射为 Query / Key / Value", {}),
        ("非局部（non-local）方式建模全局相关", {}),
        ("交叉学习：让 T1、T2 特征互相参照", {}),
        ("目标：未变化区域特征保持高相似，减少假变化", {"bold": True, "color": DARK_RED}),
    ], size=13)
    # bottom SDI full width
    add_text(s, L, Emu(3900000), FULL_W, Emu(300000),
             [("3. SDI 语义差分交互（核心创新，用图卷积）", {"size": 15, "bold": True, "color": RED})])
    add_bullets(s, L, Emu(4250000), FULL_W, Emu(2000000), [
        ("不再只对像素做差，而是把「不同层级的语义特征 / 差分特征」组织成图结构，用 GCN 做交互推理", {}),
        ("论文表述：量化 bitemporal 语义特征、其差分特征、以及二者组合之间的交互", {}),
        ("消融显示：SDI 对精度贡献最大（SECOND 上 OS +3.91%，F1scd +4.57%）", {"bold": True, "color": DARK_RED}),
        ("模块设计偏轻量，作者用 thop 统计参数与计算量增量有限", {}),
    ], size=13)
    notes(s, "TC负责稳住不变，SDI负责理清变化层次；MPANet是特征底座。强调SDI是全文最敏感模块。")

    # ========== 6 数据与设置 ==========
    s = new_slide(prs)
    add_header(s, 6, "HGINet · 实验设置")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("三、数据与实验设置", {"size": 22, "bold": True, "color": PURPLE})])
    add_text(s, L, Emu(1000000), HALF_W, Emu(300000),
             [("数据集（4 个）", {"size": 14, "bold": True, "color": RED})])
    add_bullets(s, L, Emu(1350000), HALF_W, Emu(3000000), [
        ("SECOND：公开语义变化检测基准，变化类型多、场景复杂", {}),
        ("HRSCD：城市地表语义变化数据", {}),
        ("Fuzhou（福州）：Google Earth 0.5 m，训练 6104 / 测试 2615", {}),
        ("Xiamen（厦门）：同安、翔安，训练 4542 / 测试 1946", {}),
        ("福州/厦门七类：未变化、树、耕地、建筑、水体、裸地、道路", {"size": 13}),
        ("影像切片 256×256，训练:测试 = 7:3", {"size": 13, "color": GRAY}),
    ], size=13)
    add_text(s, RIGHT_X, Emu(1000000), HALF_W, Emu(300000),
             [("训练与评价", {"size": 14, "bold": True, "color": RED})])
    add_bullets(s, RIGHT_X, Emu(1350000), HALF_W, Emu(3000000), [
        ("损失：Lt = 0.5(Ls1+Ls2) + Lc + Lsc", {"color": BLUE, "bold": True}),
        ("Lsc 用余弦约束：未变化像素拉近、变化像素推开", {}),
        ("优化：SGD，batch=8，80 epoch，RTX 3090", {}),
        ("指标：mIoU、SeK、F1scd、OS（总体分）", {"bold": True}),
        ("对比方法 9 个：BIT、GAS-Net、HRSCD-3/4、FCCDN、SSCDl、Bi-SRNet、PCFN、MTSCD-Net", {}),
        ("作者提醒：SECOND / HRSCD 存在较多错误标注", {"color": DARK_RED}),
    ], size=13)
    add_text(s, L, Emu(4600000), FULL_W, Emu(1200000),
             [("评价怎么读：SeK、F1scd 更关注「变化类」表现，避免大量未变化像素把分数抬高——这与我们关注稀有变化类是同一逻辑。",
               {"size": 14})],
             fill=LIGHT_BLUE)
    notes(s, "数据、损失、指标讲清楚，后面结果表才好读。")

    # ========== 7 主结果 ==========
    s = new_slide(prs)
    add_header(s, 7, "HGINet · 结果")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("三、主要结果：四个数据集上的最优表现", {"size": 22, "bold": True, "color": PURPLE})])
    # table
    headers = ["数据集", "F1scd", "SeK", "说明"]
    rows = [
        ["SECOND", "59.48%", "19.34%", "公开基准，变化类型复杂"],
        ["HRSCD", "64.12%", "14.55%", "城市场景语义变化"],
        ["Fuzhou", "64.45%", "18.28%", "目标小、分布不规则，更难"],
        ["Xiamen", "84.93%", "51.12%", "变化类型较少，相对更易"],
    ]
    col_w = [Emu(1600000), Emu(1600000), Emu(1600000), Emu(3700000)]
    y0 = Emu(1050000)
    x = L
    for j, h in enumerate(headers):
        add_text(s, x, y0, col_w[j], Emu(340000),
                 [(h, {"size": 13, "bold": True, "color": WHITE})],
                 fill=PURPLE, align=PP_ALIGN.CENTER)
        x = Emu(x + col_w[j])
    for i, row in enumerate(rows):
        x = L
        y = Emu(y0 + 340000 * (i + 1))
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        for j, cell in enumerate(row):
            add_text(s, x, y, col_w[j], Emu(340000),
                     [(cell, {"size": 13, "bold": j in (1, 2)})],
                     fill=bg, align=PP_ALIGN.CENTER if j < 3 else PP_ALIGN.LEFT)
            x = Emu(x + col_w[j])
    add_bullets(s, L, Emu(3000000), FULL_W, Emu(2800000), [
        ("相对 9 个基线：F1scd 在四个数据集上分别高出约 1.3–11.5 / 1.8–30.2 / 10.3–25.4 / 1.3–17.7 个百分点（论文结论）", {}),
        ("定性：假变化减少，边界更规整；季节性干扰有所缓解（作者表述）", {}),
        ("跨区域：福州训练 → 厦门测试可迁移，说明模型有一定适应性", {"color": BLUE}),
        ("难度规律：变化类型少、目标更规整时更准（厦门）；小目标、不规则分布更难（福州）", {}),
        ("与强基线 Bi-SRNet 等相比，HGINet 在复杂数据集上优势更明显", {}),
    ], size=14)
    notes(s, "读表时先说清指标含义，再报数字；强调厦门高分不代表任务简单。")

    # ========== 8 消融 ==========
    s = new_slide(prs)
    add_header(s, 8, "HGINet · 消融")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("四、消融实验：模块有没有用", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1050000), FULL_W, Emu(2000000), [
        ("在 SECOND 上逐项加入 MSCF、SFSR、TC、SDI，相对 baseline 全面提升", {}),
        ("去掉全部增强模块：OS 约降 4.91%，F1scd 约降 6.06%", {"bold": True}),
        ("单独看 SDI：OS +3.91%，F1scd +4.57% —— 贡献最大", {"bold": True, "color": DARK_RED}),
        ("MSCF 有助于大目标；SFSR 改善边界；TC 强化未变化区；SDI 改善复杂场景变化类", {}),
        ("特征提取器对比：MPANet > 纯 PVT-V2 > EfficientNet-B3 > ResNet34", {}),
        ("TC 相对传统交叉注意力（CotSR、CCA）表现更好（作者结论）", {}),
    ], size=14)
    add_text(s, L, Emu(3300000), FULL_W, Emu(280000),
             [("作者自述的局限（Discussion）", {"size": 14, "bold": True, "color": RED})])
    add_bullets(s, L, Emu(3650000), FULL_W, Emu(2000000), [
        ("边界模糊、地物被复杂背景包围时仍易出错", {}),
        ("公开数据 SECOND / HRSCD 标签噪声较多，会影响训练与评测", {"color": DARK_RED}),
        ("未来方向：边界辅助任务、位置编码、半监督伪标签", {}),
    ], size=14)
    add_text(s, L, Emu(5600000), FULL_W, Emu(1000000),
             [("读消融的方法：先看「去掉谁掉最多」，再问「为什么是它」——这里答案是：图交互（SDI）比单纯加深特征更关键。",
               {"size": 14, "bold": True})],
             fill=LIGHT_YELLOW)
    notes(s, "消融页要落到方法论判断，而不是只念数字。")

    # ========== 9 评价 ==========
    s = new_slide(prs)
    add_header(s, 9, "HGINet · 评价与启发")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("五、这篇工作好在哪、局限在哪", {"size": 22, "bold": True, "color": PURPLE})])
    add_text(s, L, Emu(1000000), HALF_W, Emu(300000),
             [("值得肯定", {"size": 15, "bold": True, "color": RED})])
    add_bullets(s, L, Emu(1350000), HALF_W, Emu(2800000), [
        ("问题抓得准：SCD 不只是「差分」，还有时序相关与层间关系", {}),
        ("结构清晰：特征（MPANet）— 不变（TC）— 变化交互（SDI）分工明确", {}),
        ("有完整实验：四数据集、九基线、消融、跨城迁移", {}),
        ("开源，可复现，适合作为 SCD 强基线", {"color": BLUE}),
        ("对未变化区的重视，符合变化检测中「背景占主导」的现实", {}),
    ], size=13)
    add_text(s, RIGHT_X, Emu(1000000), HALF_W, Emu(300000),
             [("局限与可讨论处", {"size": 15, "bold": True, "color": RED})])
    add_bullets(s, RIGHT_X, Emu(1350000), HALF_W, Emu(2800000), [
        ("任务设定仍是标准 from–to 地表分类，不是目标级「类别+状态」", {}),
        ("图建在特征层交互上，标签/目标关系结构未显式建模", {}),
        ("实例可清点性弱：输出是分割/变化图，不是目标清单", {}),
        ("公开数据标签噪声被点出，但未系统研究噪声稳健性", {}),
        ("战争/灾害等「特殊目标、极不均衡」场景未验证", {"color": DARK_RED}),
    ], size=13)
    add_text(s, L, Emu(4400000), FULL_W, Emu(2000000),
             [("组会可讨论：\n1）对变化检测来说，关系应该建在「特征」上还是「语义/标签结构」上？\n2）当变化目标稀疏且类别组合极不均衡时，图交互与样本策略谁更关键？\n3）若应用需要「哪些目标损毁成什么样」，分割式 SCD 还差哪几步？",
               {"size": 14})],
             fill=LIGHT_PURPLE)
    notes(s, "评价要公允：先肯定方法完整性，再指出任务设定与我们场景的差距。问题要能引发讨论。")

    # ========== 10 启发 ==========
    s = new_slide(prs)
    add_header(s, 10, "HGINet · 启发")
    add_text(s, L, Emu(580000), FULL_W, Emu(350000),
             [("六、对变化检测研究的几点启发", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1050000), FULL_W, Emu(3200000), [
        ("启发一：SCD 的瓶颈常在「关系」而不是「再加深网络」——时序相关、层间交互值得单独设计", {"bold": True}),
        ("启发二：图结构是表达不规则关系的有效工具；用在哪一层（特征 / 标签 / 实例）决定方法是否对症", {"bold": True}),
        ("启发三：未变化区建模不是可有可无的正则，而是降低假变化、抬高变化类指标的关键", {}),
        ("启发四：评测要用 SeK / F1scd 这类更关注变化类的指标，避免被大量背景像素误导", {}),
        ("启发五：读顶刊论文时，要同时看任务设定边界——方法强，不等于能直接回答另一个任务的问题", {"color": DARK_RED}),
        ("启发六：公开数据标签质量会系统影响结论；做自己的场景时，样本与标注分析应与方法同等重要", {}),
    ], size=14)
    add_text(s, L, Emu(4500000), FULL_W, Emu(1800000),
             [("一句话总结：\nHGINet 用「多层级特征 + 时序相关 + 图交互差分」把城市语义变化检测做扎实了；它示范了关系建模的价值，也提醒我们：换一类目标体系、换一种输出形态时，要重新问「关系建在哪一层」。",
               {"size": 15, "bold": True})],
             fill=LIGHT_YELLOW)
    notes(s, "总结给马磊听的三条：关系建模、图用在哪、指标与样本同样重要。")

    # ========== 11 文献信息 ==========
    s = new_slide(prs)
    add_header(s, 11, "HGINet · 文献信息")
    add_text(s, L, Emu(700000), FULL_W, Emu(400000),
             [("参考文献信息", {"size": 22, "bold": True, "color": PURPLE})])
    add_bullets(s, L, Emu(1300000), FULL_W, Emu(3500000), [
        ("Long, J., Li, M., Wang, X., Stein, A., 2024.", {"font": FONT_EN, "bold": True}),
        ("Semantic change detection using a hierarchical semantic graph interaction network from high-resolution remote sensing images.", {"font": FONT_EN}),
        ("ISPRS Journal of Photogrammetry and Remote Sensing, 211, 318–335.", {"font": FONT_EN}),
        ("DOI: 10.1016/j.isprsjprs.2024.04.012", {"font": FONT_EN}),
        ("代码: https://github.com/long123524/HGINet-torch", {"font": FONT_EN}),
        ("Zotero: itemKey = WY3KTIRY（题录）；PDF 附件 key = QBLZQPBN", {}),
        ("全文：Zotero storage 本地 PDF，共 18 页；本 PPT 结论均对照全文", {"color": GRAY, "size": 13}),
        ("主要对比方法：BIT, GAS-Net, HRSCD-3/4, FCCDN, SSCDl, Bi-SRNet, PCFN, MTSCD-Net", {"size": 13}),
    ], size=14)
    add_text(s, L, Emu(5000000), FULL_W, Emu(1000000),
             [("谢谢，请老师和各位同学批评指正。", {"size": 18, "bold": True, "color": PURPLE})],
             align=PP_ALIGN.CENTER)
    notes(s, "结束页；准备回答关于指标、数据集难度、与二值CD差异的问题。")

    prs.save(str(OUT))
    print("SAVED", OUT)
    print("slides", len(prs.slides))
    print("FIG_ARCH exists", FIG_ARCH.exists(), FIG_ARCH)
    # count pictures
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    n = 0
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n += 1
    print("picture_count", n)


if __name__ == "__main__":
    main()
