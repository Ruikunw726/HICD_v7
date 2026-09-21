# -*- coding: utf-8 -*-
"""组会精读：原文插图为主，圆角矩形仅用于列举，文字精简。"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(r"F:\mambacd\home\HICD_v7\docs\ISPRS2024_精读_王瑞坤.pptx")
FIG = Path(r"F:\mambacd\home\HICD_v7\docs\assets\paper_figs")
F1 = FIG / "fig1_workflow.png"
F2 = FIG / "fig2_mpanet.png"
F3 = FIG / "fig3_tc.png"
F4 = FIG / "fig4_sdi.png"
F5 = FIG / "fig5_datasets.png"
F6 = FIG / "fig6_ablation_qual.png"
T2 = FIG / "table2_ablation.png"
T3 = FIG / "table3_backbone.png"

PURPLE = RGBColor(0x6E, 0x0F, 0x6D)
RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x0D, 0x73, 0x7A)
ORANGE = RGBColor(0xC4, 0x5C, 0x26)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_PURPLE = RGBColor(0xF3, 0xE8, 0xF3)
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xF8)
LIGHT_RED = RGBColor(0xFB, 0xEB, 0xEB)
LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
LIGHT_YELLOW = RGBColor(0xFF, 0xF8, 0xE1)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)

FONT = "微软雅黑"
L = Emu(280000)
FULL_W = Emu(8500000)


def set_cjk(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    succ = {
        "a:ea": ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
        "a:cs": ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"),
    }
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.insert_element_before(el, *succ[tag])
        el.set("typeface", name)


def style(run, size=14, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_cjk(run, name)


def text(slide, x, y, w, h, segs, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    if isinstance(segs, str):
        segs = [(segs, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    for t, kw in segs:
        if t == "\n":
            p = tf.add_paragraph()
            p.alignment = kw.get("align", align)
            continue
        r = p.add_run()
        r.text = t
        style(r, size=kw.get("size", 14), bold=kw.get("bold", False),
              color=kw.get("color", INK), name=kw.get("font", FONT))
    return box


def list_card(slide, x, y, w, h, title, items, *, fill=LIGHT_BLUE, title_color=BLUE, body_size=13):
    """圆角矩形：仅用于「列举」内容"""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = title_color
    sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    text(slide, Emu(x + 100000), Emu(y + 60000), Emu(w - 200000), Emu(260000),
         [(title, {"size": 13, "bold": True, "color": title_color})])
    box = slide.shapes.add_textbox(Emu(x + 100000), Emu(y + 340000), Emu(w - 200000), Emu(h - 400000))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = "• " + it
        style(r, size=body_size, color=INK)


def header(slide, page=None):
    text(slide, L, Emu(80000), Emu(1600000), Emu(320000),
         [("论文阅读", {"size": 18, "bold": True, "color": PURPLE})])
    if page:
        text(slide, Emu(8300000), Emu(7000000), Emu(500000), Emu(220000),
             [(str(page), {"size": 10, "color": GRAY})], align=PP_ALIGN.RIGHT)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, Emu(420000), FULL_W, Emu(14000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PURPLE
    bar.line.fill.background()


def title(slide, s, y=Emu(520000)):
    text(slide, L, y, FULL_W, Emu(360000), [(s, {"size": 20, "bold": True, "color": PURPLE})])


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def img(slide, path, x, y, w, h=None):
    p = Path(path)
    if not p.exists():
        text(slide, x, y, w, Emu(300000), [(f"[缺图 {p.name}]", {"size": 12, "color": RED})])
        return
    if h is None:
        slide.shapes.add_picture(str(p), x, y, width=w)
    else:
        slide.shapes.add_picture(str(p), x, y, width=w, height=h)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ---- 1 封面 ----
    s = new(prs)
    text(s, L, Emu(1800000), FULL_W, Emu(400000),
         [("文献精读 · 组会", {"size": 16, "bold": True, "color": PURPLE})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(2300000), FULL_W, Emu(1000000),
         [("HGINet：层级语义图交互网络\n用于高分辨率遥感语义变化检测",
           {"size": 24, "bold": True})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(3600000), FULL_W, Emu(700000),
         [("Long et al., ISPRS J. Photogramm. Remote Sens., 2024\n福州大学 / 特文特大学 ITC",
           {"size": 14, "color": GRAY})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(4600000), FULL_W, Emu(400000),
         [("汇报人：王瑞坤", {"size": 16, "bold": True})], align=PP_ALIGN.CENTER)
    text(s, L, Emu(5200000), FULL_W, Emu(350000),
         [("依据 Zotero 全文 PDF；图均为论文原图", {"size": 12, "color": GRAY})], align=PP_ALIGN.CENTER)
    notes(s, "组会精读：问题—方法—结果—启发。")

    # ---- 2 问题 ----
    s = new(prs)
    header(s, 2)
    title(s, "要解决什么问题？")
    text(s, L, Emu(1000000), FULL_W, Emu(500000),
         [("语义变化检测（SCD）：既要定位「哪里变」，还要判断「变成什么类别」。",
           {"size": 16, "bold": True})])
    # 仅列举处用圆角矩形
    list_card(s, L, Emu(1700000), Emu(4000000), Emu(2200000), "难点", [
        "变化区域占比小",
        "复杂城区易误判",
        "双时相关系没建模",
    ], fill=LIGHT_RED, title_color=RED)
    list_card(s, Emu(4600000), Emu(1700000), Emu(4150000), Emu(2200000), "作者思路", [
        "特征提清楚（MPANet）",
        "稳住未变化（TC）",
        "图交互理清变化（SDI）",
    ], fill=LIGHT_GREEN, title_color=TEAL)
    text(s, L, Emu(4200000), FULL_W, Emu(800000),
         [("核心判断：", {"size": 14, "bold": True, "color": RED}),
          ("不能只做像素相减，要把双时相与多层特征之间的「关系」显式建起来。",
            {"size": 15, "bold": True})])
    notes(s, "两块列举+一句核心判断。")

    # ---- 3 总框架 原图 ----
    s = new(prs)
    header(s, 3)
    title(s, "方法总览（论文 Fig.1）")
    img(s, F1, L, Emu(950000), FULL_W, Emu(2700000))
    text(s, L, Emu(3700000), FULL_W, Emu(250000),
         [("Fig.1  HGINet 工作流：MPANet → TC → SDI → SCD 结果", {"size": 11, "color": GRAY})])
    list_card(s, L, Emu(4100000), Emu(4000000), Emu(2000000), "三步对应图中模块", [
        "(a) 提特征：MPANet",
        "(b) 时序相关：TC",
        "(c) 图交互差分：SDI",
    ], fill=LIGHT_PURPLE, title_color=PURPLE)
    text(s, Emu(4600000), Emu(4200000), Emu(4100000), Emu(1600000),
         [("输出\n\nT1/T2 语义分割图 + 变化图（from→to）",
           {"size": 15, "bold": True})])
    notes(s, "先看原图，再点三步。")

    # ---- 4 MPANet 原图 ----
    s = new(prs)
    header(s, 4)
    title(s, "模块① MPANet：特征提取（论文 Fig.2）")
    img(s, F2, L, Emu(900000), Emu(5200000), Emu(4300000))
    text(s, L, Emu(5300000), Emu(5200000), Emu(220000),
         [("Fig.2  PVT-V2 + MSCF + LAM/CBAM + SFSR", {"size": 10, "color": GRAY})])
    list_card(s, Emu(5600000), Emu(1000000), Emu(3150000), Emu(2800000), "原图在讲什么", [
        "骨干：PVT-V2 金字塔",
        "高层 F2–F4 走 MSCF",
        "低层 F1 走 CBAM",
        "SFSR 融合后 refine",
    ], fill=LIGHT_BLUE, title_color=BLUE, body_size=12)
    text(s, Emu(5600000), Emu(4100000), Emu(3150000), Emu(1800000),
         [("核心一句话\n\n把不同尺度特征提干净，后面关系才好学。",
           {"size": 14, "bold": True})])
    notes(s, "只点图上关键模块名，不讲公式。")

    # ---- 5 TC 原图 ----
    s = new(prs)
    header(s, 5)
    title(s, "模块② TC：时序相关（论文 Fig.3）")
    img(s, F3, L, Emu(950000), Emu(6200000), Emu(2700000))
    text(s, L, Emu(3750000), Emu(6200000), Emu(220000),
         [("Fig.3  双时相 Q/K/V 注意力，增强未变化区特征", {"size": 10, "color": GRAY})])
    list_card(s, L, Emu(4150000), Emu(4000000), Emu(2000000), "读图要点", [
        "T1、T2 各出 Q/V/K",
        "相似度矩阵 Softmax",
        "1×1 卷积后残差增强",
    ], fill=LIGHT_GREEN, title_color=TEAL)
    text(s, Emu(4600000), Emu(4250000), Emu(4100000), Emu(1600000),
         [("类比\n\n双时相「对答案」：该一致的尽量一致 → 假变化变少。",
           {"size": 14, "bold": True})])
    notes(s, "TC对应未变化稳住。")

    # ---- 6 SDI 原图 ----
    s = new(prs)
    header(s, 6)
    title(s, "模块③ SDI：图交互（论文 Fig.4，核心）")
    img(s, F4, L, Emu(900000), FULL_W, Emu(3400000))
    text(s, L, Emu(4350000), FULL_W, Emu(220000),
         [("Fig.4  SDI 模块结构与邻域关系（GCN 交互空间）", {"size": 10, "color": GRAY})])
    list_card(s, L, Emu(4700000), Emu(4000000), Emu(1700000), "原图在讲什么", [
        "语义特征+差分 → 图节点",
        "GCN 做层间交互",
        "再投影回变化特征",
    ], fill=LIGHT_YELLOW, title_color=ORANGE)
    list_card(s, Emu(4600000), Emu(4700000), Emu(4150000), Emu(1700000), "为何重要", [
        "消融里贡献最大",
        "复杂场景更说得清",
        "「关系」比加深网络关键",
    ], fill=LIGHT_RED, title_color=RED)
    notes(s, "全文核心：图交互。")

    # ---- 7 数据 原图 ----
    s = new(prs)
    header(s, 7)
    title(s, "数据：公开基准 + 福州/厦门（论文 Fig.5）")
    img(s, F5, L, Emu(950000), FULL_W, Emu(3400000))
    text(s, L, Emu(4400000), FULL_W, Emu(220000),
         [("Fig.5  福州（a）与厦门（b）双时相样例", {"size": 10, "color": GRAY})])
    list_card(s, L, Emu(4750000), Emu(4000000), Emu(1600000), "实验数据", [
        "SECOND、HRSCD（公开）",
        "Fuzhou、Xiamen（自建）",
        "指标：mIoU / SeK / F1scd",
    ], fill=LIGHT_BLUE, title_color=BLUE)
    text(s, Emu(4600000), Emu(4850000), Emu(4100000), Emu(1400000),
         [("SeK、F1scd 更看「变化类」，避免背景把分数抬高。",
           {"size": 14, "bold": True})])
    notes(s, "数据图+指标一句。")

    # ---- 8 结果 ----
    s = new(prs)
    header(s, 8)
    title(s, "效果如何？（主结果）")
    # bar chart with shapes
    cats = ["SECOND", "HRSCD", "Fuzhou", "Xiamen"]
    vals = [59.48, 64.12, 64.45, 84.93]
    chart_x, chart_y = L, Emu(1000000)
    chart_w, chart_h = Emu(5200000), Emu(3200000)
    # axes area
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, chart_x, chart_y, chart_w, chart_h)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xF7, 0xF7, 0xFA)
    bg.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    text(s, chart_x, Emu(chart_y - 280000), chart_w, Emu(250000),
         [("F1scd (%)，越高越好", {"size": 12, "bold": True, "color": PURPLE})])
    vmax = 100.0
    n = len(cats)
    gw = chart_w / n
    for i, (c, v) in enumerate(zip(cats, vals)):
        bh = int(chart_h * 0.85 * (v / vmax))
        bx = Emu(chart_x + i * gw + int(gw * 0.25))
        bw = Emu(int(gw * 0.5))
        by = Emu(chart_y + chart_h - 200000 - bh)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, by, bw, Emu(bh))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PURPLE
        bar.line.fill.background()
        text(s, Emu(bx - 80000), Emu(by - 220000), Emu(bw + 160000), Emu(200000),
             [(f"{v}", {"size": 11, "bold": True, "color": PURPLE})], align=PP_ALIGN.CENTER)
        text(s, Emu(chart_x + i * gw), Emu(chart_y + chart_h + 20000), Emu(gw), Emu(220000),
             [(c, {"size": 11})], align=PP_ALIGN.CENTER)

    list_card(s, Emu(5700000), Emu(1000000), Emu(3050000), Emu(2000000), "怎么读", [
        "厦门高：变化更简单",
        "公开集更难、更可信",
        "九个基线中总体最优",
    ], fill=LIGHT_GREEN, title_color=TEAL)
    list_card(s, Emu(5700000), Emu(3200000), Emu(3050000), Emu(1400000), "SeK 同步", [
        "19.3 / 14.6 / 18.3 / 51.1",
    ], fill=LIGHT_GRAY, title_color=GRAY)
    text(s, L, Emu(4800000), FULL_W, Emu(800000),
         [("结论：复杂场景提升更明显；福州训练可迁到厦门。",
           {"size": 14, "bold": True})])
    notes(s, "柱状图报F1；SeK口头补充。")

    # ---- 9 消融 原图表 ----
    s = new(prs)
    header(s, 9)
    title(s, "模块有没有用？—— 论文原表与原图")
    img(s, T2, L, Emu(900000), Emu(7200000), Emu(1900000))
    text(s, L, Emu(2850000), Emu(7200000), Emu(220000),
         [("Table 2  SECOND 消融（完整 HGINet 标红：F1scd 59.48 / SeK 19.34）", {"size": 10, "color": GRAY})])
    img(s, F6, L, Emu(3200000), Emu(7200000), Emu(2400000))
    text(s, L, Emu(5650000), Emu(7200000), Emu(220000),
         [("Fig.6  消融可视化：Baseline → +SDI/TC/… → 完整模型 vs 真值", {"size": 10, "color": GRAY})])
    text(s, Emu(7400000), Emu(3200000), Emu(1300000), Emu(2400000),
         [("SDI\n单独贡献\n最大", {"size": 13, "bold": True, "color": RED})], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    notes(s, "用论文Table2+Fig6，不自造图。")

    # ---- 10 启发 ----
    s = new(prs)
    header(s, 10)
    title(s, "启发与讨论")
    list_card(s, L, Emu(1000000), Emu(4000000), Emu(2400000), "三条启发", [
        "瓶颈常在「关系」而非加层",
        "图建在哪一层很关键",
        "先稳住未变化，变化才可信",
    ], fill=LIGHT_PURPLE, title_color=PURPLE)
    list_card(s, Emu(4600000), Emu(1000000), Emu(4150000), Emu(2400000), "作者局限", [
        "复杂边界仍易错",
        "公开数据标签噪声大",
        "偏分割式，难出实例清单",
    ], fill=LIGHT_RED, title_color=RED)
    text(s, L, Emu(3700000), FULL_W, Emu(300000),
         [("组会讨论（选 1–2 个即可）", {"size": 14, "bold": True, "color": RED})])
    list_card(s, L, Emu(4100000), FULL_W, Emu(2000000), "问题", [
        "关系建在特征上，还是建在语义/目标结构上？",
        "变化极稀疏、类别组合极不均衡时，结构和样本谁更关键？",
    ], fill=LIGHT_YELLOW, title_color=ORANGE, body_size=14)
    notes(s, "启发+局限+讨论题。")

    # ---- 11 文献 ----
    s = new(prs)
    header(s, 11)
    title(s, "文献信息")
    text(s, L, Emu(1200000), FULL_W, Emu(2000000),
         [("Long J., Li M., Wang X., Stein A., 2024.\n",
           {"size": 14, "font": "Times New Roman"}),
          ("Semantic change detection using a hierarchical semantic graph interaction network from high-resolution remote sensing images.\n",
           {"size": 14, "font": "Times New Roman"}),
          ("ISPRS Journal of Photogrammetry and Remote Sensing, 211, 318–335.\n",
           {"size": 14, "font": "Times New Roman"}),
          ("DOI: 10.1016/j.isprsjprs.2024.04.012\n",
           {"size": 13, "color": GRAY, "font": "Times New Roman"}),
          ("代码: github.com/long123524/HGINet-torch",
           {"size": 13, "color": GRAY, "font": "Times New Roman"})])
    list_card(s, L, Emu(3600000), FULL_W, Emu(1400000), "原图来源", [
        "Fig.1–6、Table2 均从 Zotero PDF 渲染页裁切",
    ], fill=LIGHT_GRAY, title_color=GRAY)
    text(s, L, Emu(5600000), FULL_W, Emu(400000),
         [("谢谢，请老师批评指正", {"size": 18, "bold": True, "color": PURPLE})],
         align=PP_ALIGN.CENTER)
    notes(s, "收尾。")

    prs.save(str(OUT))
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    n = sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    print("SAVED", OUT)
    print("slides", len(prs.slides), "pictures", n)


if __name__ == "__main__":
    main()
